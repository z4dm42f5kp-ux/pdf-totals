"""Check the profiles for other countries, languages and sectors (issue #5).

Creates small invented documents in English, German and Spanish in a temporary folder,
runs lotes_auto.py on them with each profile, and compares every batch with the
expected quantity. The default profile is checked by tools/check_examples.py.

Usage:
    python tools/check_profiles.py [--keep FOLDER]

Needs reportlab, openpyxl and pypdfium2 (requirements-dev.txt). One German document is
scanned (rendered to an image), so the OCR must be installed.

Return values: 0 when every case matches, 1 when a case differs, 2 on an error.
"""
import io
import json
import pathlib
import shutil
import subprocess
import sys
import tempfile

from openpyxl import Workbook, load_workbook
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

ROOT = pathlib.Path(__file__).resolve().parent.parent


def pdf(path, lines, scanned=False):
    """lines: [(x, y, text, size)] in points from the top-left corner."""
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter, invariant=1)
    h = letter[1]
    for x, y, text, size in lines:
        c.setFont("Helvetica-Bold" if size > 11 else "Helvetica", size)
        c.drawString(x, h - y, text)
    c.showPage()
    c.save()
    if not scanned:
        path.write_bytes(buf.getvalue())
        return
    import pypdfium2 as pdfium
    img = pdfium.PdfDocument(buf.getvalue())[0].render(scale=200 / 72).to_pil().convert("RGB")
    img.save(path, resolution=200)


def table(x0, y0, headers, rows, widths, size=9, gap=16):
    out, x = [], x0
    for hdr, w in zip(headers, widths):
        out.append((x, y0, hdr, size))
        x += w
    for i, row in enumerate(rows, 1):
        x = x0
        for val, w in zip(row, widths):
            out.append((x, y0 + i * gap, str(val), size))
            x += w
    return out


def excel(path, sheet, headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(["Recall reference list"])
    ws.append([])
    ws.append(headers)
    for r in rows:
        ws.append(list(r))
    wb.save(path)


# ------------------------------------------------------------------ cases
def case_us(d):
    """US, English, batches from the reference Excel, money 1,234.56, GS1 (10)."""
    excel(d / "recall list.xlsx", "Recall", ["Batch", "Product", "Not distributed (units)", "Withdrawn (customers)"],
          [("AB2401", "Widget", 500, 1300), ("AB2402", "Widget", 0, 300), ("CD7710", "Gadget", 250, 1200)])
    rows = [(1, "Widget 10 mg", "AB2401", "1,000", "$12.50", "$12,500.00"),
            (2, "Widget 10 mg", "AB2401", "300", "$12.50", "$3,750.00"),
            (3, "Invoice 4471", "INV-2207", "", "", "")]
    lines = [(40, 40, "RETURN AUTHORIZATION - PRODUCT RECALL", 14), (40, 60, "Date: Mar 12, 2026", 9)]
    lines += table(40, 100, ["Line", "Description", "Lot No.", "Qty (units)", "Unit Price", "Amount"], rows[:2],
                   [40, 120, 80, 80, 80, 90])
    lines += [(40, 160, "Total units returned: 1,300", 9), (40, 176, "Total amount: $16,250.00", 9)]
    pdf(d / "return authorization 001.pdf", lines)
    pdf(d / "destruction certificate.pdf", [
        (40, 40, "CERTIFICATE OF DESTRUCTION - RECALLED PRODUCT", 14),
        (40, 80, "Lot: CD7710    Quantity: 1,200 units", 10),
        (40, 100, "(01)00312345678906 (10)AB2402    Qty: 300", 10),
        (40, 130, "Destroyed on Apr 02, 2026 by an approved vendor.", 9)])
    return ["--perfil", "us"], {"ret": {"AB2401": 1300, "AB2402": 300, "CD7710": 1200},
                                "nod": {"AB2401": 500, "CD7710": 250}}


def case_de(d):
    """Germany, German, batches from the Excel, money 1.234,56 EUR, one scanned page."""
    excel(d / "Rueckrufliste.xlsx", "Rückruf", ["Charge", "Artikel", "Nicht ausgeliefert (Stück)", "Zurückgerufen (Kunden)"],
          [("K24A17", "Salbe 50 g", 120, 840), ("K24A18", "Salbe 50 g", 0, 1150), ("K24B02", "Creme", 60, 400)])
    rows = [(1, "Salbe 50 g", "K24A17", "600", "4,20 €", "2.520,00 €"),
            (2, "Salbe 50 g", "K24A18", "1.150", "4,20 €", "4.830,00 €")]
    lines = [(40, 40, "RÜCKSENDUNG - PRODUKTRÜCKRUF", 14), (40, 60, "Datum: 12.03.2026", 9)]
    lines += table(40, 100, ["Pos.", "Artikel", "Ch.-B.", "Menge (Stück)", "Einzelpreis", "Betrag"], rows,
                   [40, 120, 80, 90, 80, 90])
    lines += [(40, 160, "Gesamtmenge Stück: 1.750", 9), (40, 176, "Gesamtbetrag: 7.350,00 €", 9)]
    pdf(d / "Ruecksendung 0417.pdf", lines)
    rows2 = [(1, "Salbe 50 g", "K24A17", "240", "4,20 €", "1.008,00 €"),
             (2, "Creme 30 ml", "K24B02", "400", "3,10 €", "1.240,00 €")]
    lines2 = [(40, 40, "RETOURE APOTHEKE NORD", 14), (40, 60, "Datum: 14.03.2026", 9)]
    lines2 += table(40, 100, ["Pos.", "Artikel", "Charge", "Menge (Stück)", "Einzelpreis", "Betrag"], rows2,
                    [40, 120, 80, 90, 80, 90], size=11, gap=20)
    lines2 += [(40, 170, "Gesamtmenge Stück: 640", 11)]
    pdf(d / "Retoure Apotheke Nord (Scan).pdf", lines2, scanned=True)
    return ["--perfil", "de"], {"ret": {"K24A17": 840, "K24A18": 1150, "K24B02": 400},
                                "nod": {"K24A17": 120, "K24B02": 60}}


def case_mx_food(d):
    """Mexico, food sector, batch format from tute-lote.json, no Excel, expected totals."""
    (d / "tute-lote.json").write_text(json.dumps({
        "profile": "mx", "sector": "food", "batch_pattern": "L\\d{6}",
        "expected": {"withdrawn": 1450, "not_distributed": 300, "total": 1750}}), encoding="utf-8")
    rows = [(1, "Galleta avena 200 g", "L240311", "1,200"), (2, "Galleta avena 200 g", "L240312", "250")]
    lines = [(40, 40, "DEVOLUCION DE PRODUCTO - RETIRO DEL MERCADO", 14), (40, 60, "Fecha: 12/03/2026", 9)]
    lines += table(40, 100, ["No.", "Producto", "Lote", "Cantidad (bolsas)"], rows, [40, 150, 90, 110])
    lines += [(40, 160, "Total de bolsas devueltas: 1,450", 9)]
    pdf(d / "devolucion tienda 12.pdf", lines)
    rows2 = [(1, "Galleta avena 200 g", "L240311", "300")]
    lines2 = [(40, 40, "INVENTARIO NO DISTRIBUIDO - ALMACEN CENTRAL", 14)]
    lines2 += table(40, 90, ["No.", "Producto", "Lote", "Existencia (bolsas)"], rows2, [40, 150, 90, 110])
    pdf(d / "inventario almacen.pdf", lines2)
    return [], {"ret": {"L240311": 1200, "L240312": 250}, "nod": {"L240311": 300}}


def case_us_auto(d):
    """US, automotive sector: the VIN is the identifier (ISO 3779, 17 characters)."""
    rows = [("1HGCM82633A004352", "Accord", "1"), ("1HGCM82633A004353", "Accord", "1"),
            ("5YJ3E1EA7KF317000", "Model 3", "1")]
    lines = [(40, 40, "RECALL CAMPAIGN 26V-117 - VEHICLES RETURNED", 14)]
    lines += table(40, 90, ["VIN", "Model", "Vehicles"], rows, [170, 90, 80])
    lines += [(40, 160, "Total vehicles: 3", 9)]
    pdf(d / "campaign 26V-117 dealer 4410.pdf", lines)
    return ["--perfil", "us", "--sector", "automotive"], {
        "ret": {"1HGCM82633A004352": 1, "1HGCM82633A004353": 1, "5YJ3E1EA7KF317000": 1}, "nod": {}}


CASES = [case_us, case_de, case_mx_food, case_us_auto]


def leer_resultado(xlsx):
    ws = load_workbook(xlsx).active
    res = {"ret": {}, "nod": {}}
    fila = 3
    while True:
        lote = ws.cell(fila, 1).value
        if lote is None or str(lote).upper() in ("TOTAL", "GESAMT"):
            return res
        lote = str(lote)
        if ws.cell(fila, 2).value:
            res["ret"][lote] = ws.cell(fila, 2).value
        if ws.cell(fila, 5).value:
            res["nod"][lote] = ws.cell(fila, 5).value
        fila += 1


def main(argv):
    keep = None
    if "--keep" in argv:
        keep = pathlib.Path(argv[argv.index("--keep") + 1])
    fallas = 0
    with tempfile.TemporaryDirectory() as tmp:
        for case in CASES:
            d = pathlib.Path(tmp) / case.__name__
            d.mkdir()
            opciones, esperado = case(d)
            r = subprocess.run([sys.executable, str(ROOT / "lotes_auto.py"), str(d), "--procesos", "1"] + opciones,
                               capture_output=True, text=True)
            if keep:
                shutil.copytree(d, keep / case.__name__, dirs_exist_ok=True)
            if r.returncode != 0:
                print(f"{case.__name__}: lotes_auto.py failed with return code {r.returncode}")
                print(r.stdout[-3000:], r.stderr[-3000:])
                return 2
            obtenido = leer_resultado(d / "resultados_lotes" / "totales_lotes.xlsx")
            difs = [f"{t} {lote}: expected {esperado[t].get(lote, 0)}, got {obtenido[t].get(lote, 0)}"
                    for t in ("ret", "nod") for lote in sorted(set(esperado[t]) | set(obtenido[t]))
                    if esperado[t].get(lote, 0) != obtenido[t].get(lote, 0)]
            if case is case_mx_food and "COINCIDE" not in r.stdout:
                difs.append("the expected totals from tute-lote.json were not confirmed (no COINCIDE)")
            if difs:
                fallas += 1
                print(f"{case.__name__}: {len(difs)} difference(s)")
                for x in difs:
                    print("   " + x)
                print("   console output:\n" + "\n".join("     " + x for x in r.stdout.splitlines()[-25:]))
            else:
                n = sum(len(v) for v in esperado.values())
                print(f"{case.__name__}: OK ({n} quantities)")
    return 1 if fallas else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
