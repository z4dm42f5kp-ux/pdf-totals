"""lotes_auto.py - piezas para destruccion por numero de lote.

Lee todos los PDFs de una carpeta (devoluciones, acuses de entrega, recepciones de
producto, inventarios y reportes de retiro) y calcula por lote:

  1. Cantidad de producto retirado    piezas devueltas o recuperadas del mercado
  2. Cantidad no distribuida          piezas en inventario que nunca se distribuyeron
  3. Cantidad total para destruccion  1 + 2

Las cantidades no se buscan solo por palabras clave. Cada tabla se analiza por su
estructura:

  - de que tipo es cada celda: lote, fecha, importe, entero, codigo o texto;
  - que columna cumple  piezas x precio unitario = importe;
  - que columna suma exactamente el total impreso en el documento;
  - en los reportes de retiro, que columnas cumplen
        fabricado - distribuido = no distribuido   y
        no distribuido + retirado = total para destruccion.

Las palabras de los encabezados (PIEZAS, RECIBIDA, FACTURADA, INVENTARIO...) son
una senal mas, no la unica. Cada cantidad queda marcada como verificada o no, y al
final se comprueban los topes por lote y el total esperado. Si algo no cuadra, los
PDFs escaneados dudosos se leen otra vez con mas resolucion, y un PDF solo se marca
como MALO cuando la prueba es concluyente.
"""
import argparse
import hashlib
import json
import os
import pathlib
import re
import sys
import time
import unicodedata
from collections import Counter, defaultdict, namedtuple
from dataclasses import dataclass, field

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

VERSION = "2.3.0"

# ============================== CONFIGURACION ===============================
# Total esperado de la tabla "Cantidad total para destruccion" (None = sin comprobar).
TOTAL_ESPERADO = 344451
# Totales esperados de cada tabla (suman TOTAL_ESPERADO; None = sin comprobar).
TOTAL_RETIRADO_ESPERADO = 209908
TOTAL_NO_DISTRIBUIDA_ESPERADO = 134543
# Excel de referencia con, por lote, la cantidad no distribuida (inventarios) y la
# cantidad retirada segun clientes. Se busca solo en la carpeta y sus subcarpetas (o se
# indica con --excel). De el se toma la cantidad no distribuida; la cantidad retirada
# se calcula con los PDFs y se compara, lote por lote, con la del Excel.
EXCEL_REFERENCIA = None
# Nombre (sin extension) del Excel de referencia que se busca primero.
EXCEL_NOMBRE = "Retirado_retiro_producto_Traxicoll"
# Carpeta, dentro de la carpeta de los PDFs, donde se guardan todos los resultados.
CARPETA_RESULTADOS = "resultados_lotes"
# Con Excel de referencia, los lotes que no estan en el Excel no se suman (casi siempre
# son numeros de lote mal leidos); se listan en revisar_a_mano.txt.
SOLO_LOTES_DEL_EXCEL = True
# Topes por lote: ningun lote por si solo puede pasar de estas cantidades (no son topes
# de la suma de la tabla). Un lote que los supera indica un error de lectura o un duplicado.
MAX_RETIRADO_POR_LOTE = 48042
MAX_NO_DISTRIBUIDA_POR_LOTE = 82090
# Los reportes de retiro (fabricado / distribuido / no distribuido / retirado / total)
# dan la cantidad no distribuida de cada lote y sirven de referencia para comprobar lo
# retirado segun las devoluciones. Con True, su retirado tambien se usa para los lotes
# que no tienen ninguna devolucion en la carpeta.
USAR_REPORTE_PARA_RETIRADO = False
# Lote repetido en varias paginas del mismo PDF:
#   "iguales": se cuenta una vez si la cantidad es la misma (copia del mismo documento);
#              con cantidades distintas se suman (documentos distintos) y se avisa.
#   "todos":   cualquier repeticion se cuenta una sola vez (se conserva la mayor).
DUPLICADOS_EN_ARCHIVO = "iguales"
# PDFs cuyo nombre contiene esto se omiten (no cuentan en ningun total).
OMITIR_NOMBRE = re.compile(r"falta[\s_\-]*revisi[oó]n", re.I)
# Numero de lote valido: 8 digitos que empiezan con 25 o 26.
LOTE_ID = re.compile(r"(?<![\d.,])(2[56]\d{6})(?![\d]|[.,]\d)")
# OCR: resolucion de la primera lectura y de la segunda (solo PDFs dudosos).
OCR_DPI = 300
OCR_DPI_REINTENTO = 400
# Confianza minima del OCR para no considerar dudosa una cantidad sin verificar.
CONFIANZA_MINIMA = 0.90
# Formularios "de lado" (un producto por columna): renglon de cantidad preferido.
PRIORIDAD_RENGLON = ["devuel", "recibid", "retir", "recuper", "pieza", "cantidad", "factur"]

# ------------------------------- velocidad ----------------------------------
# Procesos que leen paginas a la vez. 0 = automatico (segun nucleos y memoria libre);
# 1 = uno solo, sin procesos paralelos. Tambien con --procesos N.
PROCESOS = 0
# Lado mayor maximo, en pixeles, de la imagen de una pagina. El OCR reduce de todos modos
# cada imagen a 2000 px; el limite solo evita imagenes gigantes (p. ej. escaneos guardados
# a 72 ppp, cuya pagina mide un metro) que antes podian tardar minutos o agotar la memoria.
LADO_MAX_PX = 6000
# Lectura 3: lado mayor maximo antes de ampliar al doble.
LADO_MAX_PX_AMPLIADA = 4400
# Lo que el OCR lee de cada pagina se guarda en resultados_lotes/cache, identificado por
# el contenido del PDF: la siguiente corrida no vuelve a leer las paginas que no cambiaron.
# Se desactiva con --sin-cache.
USAR_CACHE = True
# Cambiar este valor invalida la cache (si se cambia el OCR, la resolucion o el procesado).
CACHE_VERSION = "1"
# Aviso en pantalla cuando una pagina lleva mas de estos segundos.
AVISO_LENTO_S = 45
# Una pagina que tarda mas que esto (atascada) se deja sin leer: el PDF va a
# revisar_a_mano.txt con el motivo y el programa sigue con las demas paginas.
LIMITE_PAGINA_S = 600

# Significado de los encabezados de columna (texto sin acentos y en minusculas).
SEMANTICA = {
    "ret": r"retir|recuper|devuel|devoluc|devolv|recib|recolect|regres",
    "nod": r"no\s*distribu|sin\s*distribu|inventari|existenci|stock|almacen|retenid|cuarentena|bloquead|saldo",
    "orig": r"factur|fabric|producid|produccion|liberad|vendid|surtid|pedid|solicitad|enviad|remitid"
            r"|(?<!no )(?<!sin )distribuid",
    "caja": r"caja",
    "cant": r"cantidad|\bcant\b|pieza|\bpzas?\b|\bpzs\b|unidades|\buds\b|frasco|ampollet|ampula|vial",
    "total": r"total|destruc",
    "nocant": r"motivo|folio|clave|codigo|telefono|\bcp\b|\bsap\b",
}
SEMANTICA = {k: re.compile(v) for k, v in SEMANTICA.items()}

# Tipo de documento por pagina: (expresion, peso).
TIPO_DOCUMENTO = {
    "ret": [(r"devoluci", 3), (r"devuelt", 2), (r"retiro", 3), (r"retirad", 3), (r"recall", 3),
            (r"recuperad", 2), (r"recolecci", 2), (r"acuse de entrega", 2),
            (r"recepcion de producto", 2), (r"anomalia", 2), (r"motivo", 1), (r"faltante", 1),
            (r"maltratad", 1)],
    "nod": [(r"inventario", 3), (r"existencia", 3), (r"stock", 2), (r"no distribuid", 4),
            (r"sin distribuir", 4), (r"cuarentena", 2), (r"retenid", 2), (r"bloquead", 2)],
}
# ============================================================================

TIPOS = {"ret": "retirado", "nod": "no distribuida"}
MESES = r"(?:ene|feb|mar|abr|may|jun|jul|ago|sep|set|oct|nov|dic)[a-z]*\.?"
FECHA_RE = re.compile(
    r"^(?:\d{1,2}[/\-.]\d{1,2}[/\-.]\d{2,4}|\d{4}[/\-.]\d{1,2}[/\-.]\d{1,2}"
    r"|\d{1,2}[/\-. ]?" + MESES + r"[/\-. ]?\d{2,4}|" + MESES + r"[/\-. ]?\d{2,4}|\d{1,2}/\d{4})$", re.I)
DINERO_RE = re.compile(r"^(?:[$sS]\s*)?(\d{1,3}(?:,\d{3})+|\d+)\.\d{2}$|^[$sS]\s*(\d{1,3}(?:,\d{3})+|\d+)$")
UNIDAD_RE = re.compile(r"^(\d[\d.,]*)\s*(pzas?|piezas?|pzs|pz|uds?|unidades|frascos?|amp\w*|viales?|cajas?)\.?$", re.I)
MARCAS = "✓✔√/\\|*·~'\"`´¨°º-–—_•▪◂▸►"
LOTE_PREFIJO_OK = re.compile(r"(?:l[o0]te|lot|serie|no|num|n|l)$", re.I)
LOTE_PALABRA = re.compile(r"l[o0]te|serie|serial|n/s|s/n", re.I)

# Etiquetas en lineas de texto ("LOTE: 25240310" / "CANTIDAD EN PIEZAS: 1,250").
CANT_PALABRAS = r"cantidad|cant|piezas|pieza|pzas|pza|pzs|pz|unidades|unid|uds"
Q = r"(?:" + CANT_PALABRAS + r")(?![a-z])\.?"
CANT_ETIQUETA = re.compile(
    Q + r"(?:\s*(?:/|en|de)?\s*" + Q + r")*\s*(?:\(\w*\))?\s*[:=\-]?\s*(?P<n>\d[\d.,]*)", re.I)
LOTE_ETIQUETA = re.compile(r"(?:l[o0]te|serial|serie|n/s|s/n)\.?\s*(?:n[º°o]\.?|n[uú]m\.?|n[uú]mero)?"
                           r"\s*[:#\-]?\s*(?P<id>\d{8})", re.I)

# Valores del perfil predeterminado (este retiro) que el resto del programa consulta.
LOTE_MAYUSCULAS = False                   # True: los lotes se comparan en mayusculas
FILA_TOTAL_RE = re.compile(r"pieza|cantidad|pzas|unidades")   # renglon "TOTAL ... PIEZAS"
TOTAL_PALABRA_RE = re.compile(r"total")
PIEZA_RE = re.compile(r"pieza")           # "cajas" que en realidad cuentan piezas
CAJA_PREFIJOS = ("caj",)
EXCEL_NOD_RE = re.compile(r"no\s*distribu")
EXCEL_RET_RE = re.compile(r"retirad")
EXCEL_RET_EXCLUYE = re.compile(r"total|destruc")
EXCEL_LOTE_RE = None                      # perfiles sin patron de lote: columna de lotes del Excel
NUMEROS = "mx"                            # "mx": 1,234.56   "auto": tambien 1.234,56
TEXTOS_EXCEL = {"hoja": "Lotes", "titulos": ["Cantidad de producto retirado", "Cantidad no distribuida",
                                             "Cantidad total para destrucción"],
                "lote": "Lote", "cantidad": "Cantidad en Piezas", "total": "TOTAL",
                "esperado": "Esperado", "diferencia": "Diferencia"}


# ================================ PERFILES ==================================
# Un perfil adapta el programa a otro pais, idioma o sector (issue #5). El perfil
# predeterminado, "traxicoll", es exactamente el comportamiento de siempre: no cambia
# nada de lo anterior. Los demas perfiles agregan vocabulario en ingles, aleman, frances
# e italiano, aceptan importes con coma decimal (1.234,56 EUR), no comprueban totales
# esperados ni topes (no se conocen) y toman el formato del numero de lote de la
# configuracion, del sector o de la columna de lotes del Excel de referencia.
# Fuentes del vocabulario: GS1 (AI 10 lote, 21 serie, 17 caducidad), AMG §10 ("Ch.-B.",
# "verwendbar bis"), LKV §1 (lote de alimentos con "L"), reglamento GPSR de la UE
# (type, batch or serial number), formularios de retiro de la FDA y del USDA (Lot Code,
# Quantity on Hand, Destroyed), ISO 3779 (VIN de 17 caracteres).
PERFIL_PREDETERMINADO = "traxicoll"
CONFIG_NOMBRE = "tute-lote.json"

VOCABULARIO = {
    # Expresiones sobre texto sin acentos y en minusculas (sin_acentos), salvo las que
    # dicen "texto original", que se buscan en el texto tal como se leyo (re.I).
    "en": {
        "ret": r"return|recover|recall|withdraw|received|collected",
        "nod": r"not\s*distribut|undistribut|inventory|on\s*hand|warehouse|quarantin|\bheld\b|blocked",
        "orig": r"invoic|shipped|\bsold\b|ordered|produced|manufactur|released|(?<!not )(?<!un)distributed",
        "caja": r"\bcases?\b|\bbox(?:es)?\b|carton|\bctn\b",
        "cant": r"quantity|\bqty\b|\bunits?\b|\bpcs\b|pieces?|\beach\b|\bea\b|\bcount\b",
        "total": r"total|destroy|destruct",
        "nocant": r"reason|invoice\s*(?:no|#|number)|\bcode\b|\bsku\b|\bndc\b|\bgtin\b|\bupc\b|\bitem\b|phone|\bzip\b",
        "doc_ret": [(r"return", 3), (r"recall", 3), (r"withdraw", 3), (r"recovered", 2),
                    (r"proof of delivery", 2), (r"goods received", 2), (r"destruction", 1)],
        "doc_nod": [(r"inventory", 3), (r"not distributed", 4), (r"undistributed", 4), (r"on hand", 2),
                    (r"quarantine", 2), (r"warehouse", 2)],
        "meses": r"jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec",
        "cant_palabras": r"quantity|qty|units|unit|pcs|pieces|piece|pc|each|ea|count",   # texto original
        "lote_etiquetas": r"\blot\s*(?:code|number|no|#)?|\bbatch(?:\s*(?:code|number|no|#))?|\bserial|\bs/n|\(10\)",
        "lote_palabras": r"\blot\b|batch|serial|s/n",
        "unidades": r"pcs|pc|pieces?|units?|ea|each|cases?|box(?:es)?|ctns?|cartons?",
        "cajas": ("case", "box", "carton", "ctn"),
        "prioridad": ["returned", "received", "withdrawn", "recovered", "pieces", "quantity", "invoiced"],
        "excel_nod": r"not\s*distribut|undistribut",
        "excel_ret": r"withdrawn|returned|recalled|recovered",
        "excel_excluye": r"total|destroy|destruct",
        "excel_lote": r"\blot\b|lot\s*(?:no|#|number|code)|batch|serial",
        "omitir": r"to[\s_\-]*review|needs[\s_\-]*review",
    },
    "de": {
        "ret": r"zuruck|retour|rucksend|ruckgabe|ruckruf|erhalten|eingegangen",
        "nod": r"nicht\s*(?:ausgeliefert|vertrieben|verteilt)|lagerbestand|\bbestand|\blager\b|quarantan|gesperrt",
        "orig": r"berechnet|fakturiert|(?<!nicht )geliefert|verkauft|bestellt|produziert|hergestellt|freigegeben",
        "caja": r"karton|kiste",
        "cant": r"menge|anzahl|stuck|\bstk\b|\bst\b",
        "total": r"gesamt|summe|vernicht",
        "nocant": r"grund|rechnungs|artikel|\bpzn\b|\bean\b|telefon|\bplz\b",
        "doc_ret": [(r"rucksendung", 3), (r"retoure", 3), (r"ruckruf", 3), (r"ruckgabe", 2),
                    (r"zuruckgesandt", 2), (r"wareneingang", 2), (r"vernichtung", 1)],
        "doc_nod": [(r"lagerbestand", 3), (r"nicht ausgeliefert", 4), (r"nicht vertrieben", 4),
                    (r"bestand", 2), (r"quarantane", 2), (r"gesperrt", 2)],
        "meses": r"jan|feb|m[aä]r|apr|mai|jun|jul|aug|sep|okt|nov|dez",
        "cant_palabras": r"menge|anzahl|stück|stueck|stuck|stk|st",
        "lote_etiquetas": r"\bch\.?\s*-?\s*b|\bchargen?(?:\s*-?\s*(?:nr|nummer|bezeichnung))?|\blos(?:\s*-?\s*(?:nr|nummer))?|\bfin\b",
        "lote_palabras": r"charge|ch\.?\s*-?\s*b\b|\blos\b",
        "unidades": r"stk|stück|stueck|stuck|st|kartons?|kisten?",
        "cajas": ("karton", "kiste"),
        "prioridad": ["zuruck", "erhalten", "ruckgabe", "menge", "stuck"],
        "excel_nod": r"nicht\s*(?:ausgeliefert|vertrieben|verteilt)",
        "excel_ret": r"zuruck|ruckruf|retour",
        "excel_excluye": r"gesamt|summe|vernicht",
        "excel_lote": r"charge|ch\.?\s*-?\s*b|\blos\b",
        "omitir": r"zu[\s_\-]*pr(?:ü|ue|u)fen|nachpr(?:ü|ue|u)fen",
    },
    "fr": {
        "ret": r"retour|rappel|retir|recu\b|recupere",
        "nod": r"non\s*distribu|inventaire|en\s*stock|entrepot",
        "orig": r"facture|expedie|vendu|commande|fabrique|libere",
        "caja": r"carton|caisse",
        "cant": r"quantite|\bqte\b|unites?",
        "total": r"total|destruct",
        "nocant": r"motif|facture\s*n|\bcode\b|article",
        "doc_ret": [(r"retour", 3), (r"rappel", 3), (r"retrait", 3)],
        "doc_nod": [(r"inventaire", 3), (r"non distribu", 4)],
        "meses": r"janv|fevr|févr|mars|avr|mai|juin|juil|aout|août|sept|oct|nov|dec|déc",
        "cant_palabras": r"quantité|quantite|qté|qte|unités|unites",
        "lote_etiquetas": r"\bn[°o]\s*de\s*lot|\blot\b",
        "lote_palabras": r"\blot\b",
        "unidades": r"unités|unites|pièces|pieces|cartons?",
        "cajas": ("carton", "caisse"),
        "prioridad": ["retourn", "recu", "quantite"],
        "excel_nod": r"non\s*distribu",
        "excel_ret": r"retire|retourne|rappel",
        "excel_excluye": r"total|destruct",
        "excel_lote": r"\blot\b",
        "omitir": r"a[\s_\-]*verifier",
    },
    "it": {
        "ret": r"\bres[oi]\b|richiam|ritirat|ricevut|recuperat",
        "nod": r"non\s*distribu|giacenz|magazzino",
        "orig": r"fatturat|spedit|vendut|ordinat|prodott",
        "caja": r"cartone|scatol|cass[ae]",
        "cant": r"quantita|\bqta\b|pezzi",
        "total": r"totale|distruz",
        "nocant": r"motivo|fattura|codice|articolo",
        "doc_ret": [(r"reso", 3), (r"richiamo", 3), (r"ritiro", 3)],
        "doc_nod": [(r"giacenza", 3), (r"non distribu", 4)],
        "meses": r"gen|feb|mar|apr|mag|giu|lug|ago|set|ott|nov|dic",
        "cant_palabras": r"quantità|quantita|qtà|qta|pezzi",
        "lote_etiquetas": r"\blotto",
        "lote_palabras": r"lotto",
        "unidades": r"pezzi|pz|cartoni?",
        "cajas": ("carton", "scatol"),
        "prioridad": ["res", "ricevut", "pezzi", "quantita"],
        "excel_nod": r"non\s*distribu",
        "excel_ret": r"ritirat|res[oi]\b|richiam",
        "excel_excluye": r"totale|distruz",
        "excel_lote": r"lotto",
        "omitir": r"da[\s_\-]*verificare",
    },
}

# Sectores: unidades de conteo y etiquetas de identificacion propias del sector.
SECTORES = {
    "pharma": {"unidades": r"frascos?|ampolletas?|viales?|vials?|ampoules?|bottles?|blisters?|packs?"
                           r"|flaschen?|ampullen?|packungen?|durchstechflaschen?",
               "cant": r"frasco|ampollet|vial|ampoule|bottle|blister|\bpack|flasche|ampulle|packung",
               "lote_etiquetas": r"\(10\)|\bch\.?\s*-?\s*b"},
    "devices": {"unidades": r"devices?|kits?|units?|unidades|equipos?|ger(?:ä|ae|a)te?",
                "cant": r"device|\bkit|equipo|gerat",
                "lote_etiquetas": r"\(10\)|\(21\)|\bref\b|\budi\b|\bsn\b"},
    "food": {"unidades": r"bolsas?|latas?|botellas?|paquetes?|frascos?|bags?|cans?|bottles?|jars?|packs?"
                         r"|beutel|dosen?|flaschen?|packungen?|gl(?:ä|ae|a)ser",
             "cant": r"bolsa|lata|botella|paquete|\bbag|\bcan\b|\bcans\b|bottle|\bjar|beutel|dose|flasche|packung",
             "lote_etiquetas": r"\bl\b|lot\s*code|\bmhd\b|best\s*before|consumir\s*antes"},
    "cosmetics": {"unidades": r"tubos?|frascos?|tubes?|jars?|bottles?|tuben?|tiegel|flaschen?",
                  "cant": r"tubo|tube|\bjar|bottle|tiegel|flasche",
                  "lote_etiquetas": r"\bbatch\s*code|\blot\b"},
    "automotive": {"unidades": r"veh[ií]culos?|vehicles?|fahrzeuge?|unidades|units?",
                   "cant": r"vehiculo|vehicle|fahrzeug",
                   "lote_etiquetas": r"\bvin\b|\bniv\b|\bfin\b|fahrgestell",
                   # ISO 3779: 17 caracteres sin I, O ni Q
                   "lote": r"[A-HJ-NPR-Z0-9]{17}"},
    "consumer": {"unidades": r"art[ií]culos?|items?|art?ikel|units?|unidades",
                 "cant": r"articulo|\bitem|artikel",
                 "lote_etiquetas": r"\bmodel|\bmodelo|\bdate\s*code"},
    "chemicals": {"unidades": r"tambor(?:es)?|garrafas?|sacos?|drums?|pails?|bags?|containers?|f(?:ä|ae|a)sser|fass"
                              r"|kanister|s(?:ä|ae|a)cke?|gebinde",
                  "cant": r"tambor|garrafa|saco|drum|\bpail|container|fass|kanister|gebinde",
                  "lote_etiquetas": r"\bbatch\b"},
    "general": {},
}

PERFILES = {
    # nombre: idiomas (el espanol siempre se incluye), formato de numeros, sector
    PERFIL_PREDETERMINADO: {"idiomas": [], "numeros": "mx", "sector": "pharma"},
    "mx": {"idiomas": ["en"], "numeros": "mx", "sector": "general", "textos": "es"},
    "us": {"idiomas": ["en"], "numeros": "auto", "sector": "general", "textos": "en"},
    "de": {"idiomas": ["de", "en"], "numeros": "auto", "sector": "general", "textos": "de"},
    "eu": {"idiomas": ["en", "de", "fr", "it"], "numeros": "auto", "sector": "general", "textos": "en"},
}

TEXTOS_EXCEL_IDIOMA = {
    "en": {"hoja": "Batches", "titulos": ["Quantity withdrawn", "Quantity not distributed",
                                          "Total quantity for destruction"],
           "lote": "Batch", "cantidad": "Quantity (pieces)", "total": "TOTAL",
           "esperado": "Expected", "diferencia": "Difference"},
    "de": {"hoja": "Chargen", "titulos": ["Zurückgerufene Menge", "Nicht ausgelieferte Menge",
                                          "Gesamtmenge zur Vernichtung"],
           "lote": "Charge", "cantidad": "Menge (Stück)", "total": "GESAMT",
           "esperado": "Erwartet", "diferencia": "Differenz"},
}


class ConfigNoValida(Exception):
    """La configuracion (perfil, sector o tute-lote.json) no es valida."""


def _une(base, extras, flags=0):
    """Regex compilada con la expresion base y las expresiones extra como alternativas."""
    partes = [p for p in [base] + list(extras) if p]
    return re.compile("|".join(f"(?:{p})" for p in partes), flags)


def _patron_lote(nucleo):
    """Expresion de lote completa a partir del patron del lote en si (grupo 1 = lote)."""
    return re.compile(r"(?<![A-Za-z0-9.,])(" + nucleo + r")(?![A-Za-z0-9]|[.,]\d)",
                      re.I if LOTE_MAYUSCULAS else 0)


def poner_patron_lote(nucleo):
    """Cambia el formato del numero de lote (perfiles distintos del predeterminado)."""
    global LOTE_ID, LOTE_ETIQUETA
    try:
        LOTE_ID = _patron_lote(nucleo)
    except re.error as e:
        raise ConfigNoValida(f"el patron de lote {nucleo!r} no es una expresion valida: {e}")
    LOTE_ETIQUETA = re.compile(r"(?:" + _ETIQUETAS_LOTE + r")\.?\s*(?:n[º°o]\.?|n[uú]m\.?|n[uú]mero|nr\.?|no\.?|#)?"
                               r"\s*[:#\-]?\s*(?P<id>" + nucleo + r")", re.I)


_ETIQUETAS_LOTE = r"l[o0]te|serial|serie|n/s|s/n"


def aplicar_perfil(nombre=None, sector=None, config=None):
    """Aplica un perfil, un sector y la configuracion de tute-lote.json. Con el perfil
    predeterminado, sin sector distinto y sin configuracion no cambia nada.
    Devuelve una lista de lineas que describen lo aplicado (para la consola)."""
    global SEMANTICA, TIPO_DOCUMENTO, MESES, FECHA_RE, DINERO_RE, UNIDAD_RE, LOTE_PREFIJO_OK, LOTE_PALABRA
    global CANT_PALABRAS, Q, CANT_ETIQUETA, PRIORIDAD_RENGLON, FILA_TOTAL_RE, TOTAL_PALABRA_RE, PIEZA_RE
    global CAJA_PREFIJOS, EXCEL_NOD_RE, EXCEL_RET_RE, EXCEL_RET_EXCLUYE, EXCEL_LOTE_RE, NUMEROS, TEXTOS_EXCEL
    global TOTAL_ESPERADO, TOTAL_RETIRADO_ESPERADO, TOTAL_NO_DISTRIBUIDA_ESPERADO, EXCEL_NOMBRE
    global SOLO_LOTES_DEL_EXCEL, MAX_RETIRADO_POR_LOTE, MAX_NO_DISTRIBUIDA_POR_LOTE, OMITIR_NOMBRE
    global LOTE_MAYUSCULAS, LOTE_ID, LOTE_ETIQUETA, _ETIQUETAS_LOTE, RELEVANTE_RE, CASI_LOTE_RE
    config = dict(config or {})
    nombre = nombre or config.pop("profile", None) or PERFIL_PREDETERMINADO
    sector = sector or config.pop("sector", None)
    config.pop("profile", None)
    config.pop("sector", None)
    if nombre not in PERFILES:
        raise ConfigNoValida(f"perfil desconocido {nombre!r}; perfiles: {', '.join(PERFILES)}")
    perfil = PERFILES[nombre]
    sector = sector or perfil["sector"]
    if sector not in SECTORES:
        raise ConfigNoValida(f"sector desconocido {sector!r}; sectores: {', '.join(SECTORES)}")
    desc = []
    extra_idiomas = config.pop("languages", None) or []
    formato = config.pop("number_format", None)
    salida = config.pop("output_language", None)
    if not isinstance(extra_idiomas, list) or any(i not in VOCABULARIO and i != "es" for i in extra_idiomas):
        raise ConfigNoValida(f"'languages' debe ser una lista con: es, {', '.join(VOCABULARIO)}")
    if formato not in (None, "mx", "auto"):
        raise ConfigNoValida("'number_format' debe ser \"mx\" (1,234.56) o \"auto\" (tambien 1.234,56)")
    if salida not in (None, "es", *TEXTOS_EXCEL_IDIOMA):
        raise ConfigNoValida(f"'output_language' debe ser: es, {', '.join(TEXTOS_EXCEL_IDIOMA)}")
    predeterminado = (nombre == PERFIL_PREDETERMINADO and sector == perfil["sector"] and not extra_idiomas
                      and formato in (None, perfil["numeros"]) and salida in (None, "es"))
    if not predeterminado:
        idiomas = list(dict.fromkeys(i for i in perfil["idiomas"] + extra_idiomas if i in VOCABULARIO))
        voc = [VOCABULARIO[i] for i in idiomas]
        sec = SECTORES[sector]
        k = lambda clave: [v[clave] for v in voc if v.get(clave)]
        SEMANTICA = {c: _une(SEMANTICA[c].pattern, k(c) + ([sec["cant"]] if c == "cant" and sec.get("cant") else []))
                     for c in SEMANTICA}
        TIPO_DOCUMENTO = {c: TIPO_DOCUMENTO[c] + [x for v in voc for x in v["doc_" + c]] for c in TIPO_DOCUMENTO}
        MESES = r"(?:" + "|".join([r"ene|feb|mar|abr|may|jun|jul|ago|sep|set|oct|nov|dic"] + k("meses")) + r")[a-z]*\.?"
        FECHA_RE = re.compile(
            r"^(?:\d{1,2}[/\-.]\d{1,2}[/\-.]\d{2,4}|\d{4}[/\-.]\d{1,2}[/\-.]\d{1,2}"
            r"|\d{1,2}[/\-. ]?" + MESES + r"[/\-. ]?\d{2,4}|" + MESES + r"[/\-. ]?\d{2,4}|\d{1,2}/\d{4})$", re.I)
        unidades = "|".join([r"pzas?|piezas?|pzs|pz|uds?|unidades|frascos?|amp\w*|viales?|cajas?"] + k("unidades")
                            + ([sec["unidades"]] if sec.get("unidades") else []))
        UNIDAD_RE = re.compile(r"^(\d[\d.,]*)\s*(" + unidades + r")\.?$", re.I)
        CAJA_PREFIJOS = tuple(dict.fromkeys(("caj",) + tuple(x for v in voc for x in v["cajas"])))
        LOTE_PREFIJO_OK = re.compile(r"(?:l[o0]te|lot|serie|no|num|n|l|batch|charge|chb|ch|los|lotto|nr|sn)$", re.I)
        LOTE_PALABRA = _une(LOTE_PALABRA.pattern, k("lote_palabras"), re.I)
        CANT_PALABRAS = "|".join([CANT_PALABRAS] + k("cant_palabras"))
        Q = r"(?:" + CANT_PALABRAS + r")(?![a-z])\.?"
        CANT_ETIQUETA = re.compile(
            Q + r"(?:\s*(?:/|en|de|in|of)?\s*" + Q + r")*\s*(?:\(\w*\))?\s*[:=\-]?\s*(?P<n>\d[\d.,]*)", re.I)
        _ETIQUETAS_LOTE = "|".join([_ETIQUETAS_LOTE] + k("lote_etiquetas")
                                   + ([sec["lote_etiquetas"]] if sec.get("lote_etiquetas") else []))
        PRIORIDAD_RENGLON = PRIORIDAD_RENGLON + [p for v in voc for p in v["prioridad"]]
        FILA_TOTAL_RE = _une(FILA_TOTAL_RE.pattern, k("cant") + ([sec["cant"]] if sec.get("cant") else []))
        TOTAL_PALABRA_RE = _une(TOTAL_PALABRA_RE.pattern, k("total"))
        PIEZA_RE = _une(PIEZA_RE.pattern, [r"piece|\bpcs\b|stuck|\bstk\b|pezzi|unit"])
        EXCEL_NOD_RE = _une(EXCEL_NOD_RE.pattern, k("excel_nod"))
        EXCEL_RET_RE = _une(EXCEL_RET_RE.pattern, k("excel_ret"))
        EXCEL_RET_EXCLUYE = _une(EXCEL_RET_EXCLUYE.pattern, k("excel_excluye"))
        OMITIR_NOMBRE = _une(OMITIR_NOMBRE.pattern, k("omitir"), re.I)
        NUMEROS = formato or perfil["numeros"]
        if NUMEROS == "auto":
            DINERO_RE = re.compile(
                r"^(?:[$sS€]\s*)?(?:\d{1,3}(?:,\d{3})+|\d+)\.\d{2}(?:\s*(?:usd|mxn|eur|€|\$))?$"
                r"|^(?:€\s*)?(?:\d{1,3}(?:\.\d{3})+|\d+),\d{2}(?:\s*(?:eur|€|chf|fr\.?))?$"
                r"|^[$€]\s*(?:\d{1,3}(?:[,.]\d{3})+|\d+)$", re.I)
        TEXTOS_EXCEL = TEXTOS_EXCEL_IDIOMA.get(salida or perfil.get("textos"), TEXTOS_EXCEL)
        if nombre != PERFIL_PREDETERMINADO:
            # totales, topes, Excel y lote de este retiro: no valen para otros
            TOTAL_ESPERADO = TOTAL_RETIRADO_ESPERADO = TOTAL_NO_DISTRIBUIDA_ESPERADO = None
            MAX_RETIRADO_POR_LOTE = MAX_NO_DISTRIBUIDA_POR_LOTE = float("inf")
            EXCEL_NOMBRE = None
            LOTE_MAYUSCULAS = True
            EXCEL_LOTE_RE = _une("", [r"l[o0]te"] + k("excel_lote"))
            LOTE_ID = None                 # se toma del sector, de la configuracion o del Excel
        if sec.get("lote"):
            poner_patron_lote(sec["lote"])
        elif LOTE_ID is not None:          # perfil predeterminado con otro sector: mismo lote
            LOTE_ETIQUETA = re.compile(r"(?:" + _ETIQUETAS_LOTE + r")\.?\s*(?:n[º°o]\.?|n[uú]m\.?|n[uú]mero)?"
                                       r"\s*[:#\-]?\s*(?P<id>\d{8})", re.I)
        RELEVANTE_RE = _une(RELEVANTE_RE.pattern, k("lote_palabras") + k("cant") + k("ret"), re.I)
        CASI_LOTE_RE = CASI_LOTE_RE if nombre == PERFIL_PREDETERMINADO else None
        desc.append(f"perfil {nombre}, sector {sector}"
                    + (f", idiomas: espanol, {', '.join(idiomas)}" if idiomas else ""))
    # ---- tute-lote.json ----
    if "batch_pattern" in config:
        poner_patron_lote(str(config.pop("batch_pattern")))
        desc.append("formato de lote de la configuracion")
    esperado = config.pop("expected", None)
    if esperado is not None:
        if not isinstance(esperado, dict):
            raise ConfigNoValida("'expected' debe ser un objeto con withdrawn, not_distributed y total")
        TOTAL_RETIRADO_ESPERADO = _entero_o_none(esperado.get("withdrawn"), "expected.withdrawn")
        TOTAL_NO_DISTRIBUIDA_ESPERADO = _entero_o_none(esperado.get("not_distributed"), "expected.not_distributed")
        TOTAL_ESPERADO = _entero_o_none(esperado.get("total"), "expected.total")
    topes = config.pop("max_per_batch", None)
    if topes is not None:
        if not isinstance(topes, dict):
            raise ConfigNoValida("'max_per_batch' debe ser un objeto con withdrawn y not_distributed")
        a = _entero_o_none(topes.get("withdrawn"), "max_per_batch.withdrawn")
        b = _entero_o_none(topes.get("not_distributed"), "max_per_batch.not_distributed")
        MAX_RETIRADO_POR_LOTE = float("inf") if a is None else a
        MAX_NO_DISTRIBUIDA_POR_LOTE = float("inf") if b is None else b
    if "reference_excel" in config:
        EXCEL_NOMBRE = config.pop("reference_excel") or None
    if "only_batches_in_excel" in config:
        SOLO_LOTES_DEL_EXCEL = bool(config.pop("only_batches_in_excel"))
    if "skip_name_pattern" in config:
        try:
            OMITIR_NOMBRE = re.compile(str(config.pop("skip_name_pattern")), re.I)
        except re.error as e:
            raise ConfigNoValida(f"skip_name_pattern no es una expresion valida: {e}")
    if config:
        raise ConfigNoValida(f"clave(s) desconocida(s) en la configuracion: {', '.join(sorted(config))}")
    return desc


def _entero_o_none(v, clave):
    if v is None or v == 0:
        return None
    if isinstance(v, bool) or not isinstance(v, int) or v < 0:
        raise ConfigNoValida(f"'{clave}' debe ser un numero entero (o null)")
    return v


def leer_config(ruta):
    try:
        datos = json.loads(pathlib.Path(ruta).read_text(encoding="utf-8-sig"))
    except (OSError, ValueError) as e:
        raise ConfigNoValida(f"no se pudo leer {ruta}: {e}")
    if not isinstance(datos, dict):
        raise ConfigNoValida(f"{ruta} debe contener un objeto JSON")
    return datos


# ------------------------------- utilidades --------------------------------
def sin_acentos(texto):
    t = unicodedata.normalize("NFKD", texto or "")
    return "".join(c for c in t if not unicodedata.combining(c)).lower()


def lotes_en(texto):
    """Todos los lotes validos del texto, en orden."""
    res = []
    if LOTE_ID is None:                   # perfil sin formato de lote todavia (antes del Excel)
        return res
    for m in LOTE_ID.finditer(texto or ""):
        previo = re.search(r"[A-Za-z]+$", texto[:m.start()])
        if previo and not LOTE_PREFIJO_OK.search(previo.group()):
            continue                      # p. ej. una factura FE25012345
        res.append(m.group(1).upper() if LOTE_MAYUSCULAS else m.group(1))
    return res


def valid_lote(texto):
    """Primer lote valido del texto o None (se conserva el nombre por compatibilidad)."""
    lotes = lotes_en(texto)
    return lotes[0] if lotes else None


def a_entero(s):
    """Entero en notacion mexicana: 1,250 / 1.250 (coma leida como punto) / 1250."""
    s = (s or "").replace(" ", "")
    if re.fullmatch(r"\d{1,3}(?:[,.]\d{3})+", s):
        return int(re.sub(r"[,.]", "", s))
    if re.fullmatch(r"\d+", s):
        return int(s)
    return None


def a_dinero(s):
    if NUMEROS == "auto":
        s = re.sub(r"(?i)^[$sS€]\s*|\s*(?:usd|mxn|eur|€|\$|chf|fr\.?)$", "", (s or "").strip())
        if re.fullmatch(r"(?:\d{1,3}(?:\.\d{3})+|\d+),\d{2}", s):          # 1.234,56
            s = s.replace(".", "").replace(",", ".")
        try:
            return float(s.replace(",", ""))
        except ValueError:
            return None
    s = re.sub(r"^[$sS]\s*", "", (s or "").strip()).replace(",", "")
    try:
        return float(s)
    except ValueError:
        return None


def fmt(v):
    return f"{int(v):,}" if float(v).is_integer() else f"{v:,.2f}"


def orden_lote(s):
    return [int(t) if t.isdigit() else t for t in re.split(r"(\d+)", str(s))]


def semantica(texto):
    t = sin_acentos(texto)
    return {k: bool(rx.search(t)) for k, rx in SEMANTICA.items()}


def tipo_documento(texto):
    """'ret', 'nod' o None segun el vocabulario de la pagina."""
    t = sin_acentos(texto)
    puntos = {k: sum(p for rx, p in reglas if re.search(rx, t)) for k, reglas in TIPO_DOCUMENTO.items()}
    if puntos["ret"] >= puntos["nod"] + 2:
        return "ret"
    if puntos["nod"] >= puntos["ret"] + 2:
        return "nod"
    return None


# -------------------------------- modelo -----------------------------------
Tok = namedtuple("Tok", "x0 x1 y h text conf caja")
Celda = namedtuple("Celda", "x0 x1 text conf tipo valor")
Item = namedtuple("Item", "fila lote celda_lote celdas")


class Renglon:
    def __init__(self, y, h):
        self.y, self.h, self.toks, self.texto = y, h, [], ""


@dataclass
class Entrada:
    archivo: str
    pagina: int
    lote: str
    piezas: int
    tipo: str                 # "ret" o "nod"
    forma: str                # "tabla", "de lado", "etiquetas"
    conf: float = 1.0
    verificado: list = field(default_factory=list)
    fallas: list = field(default_factory=list)
    notas: list = field(default_factory=list)
    excluido: str = ""

    def clave(self):
        return (self.tipo, self.lote, self.piezas)


@dataclass
class Referencia:
    archivo: str
    pagina: int
    lote: str
    ret: object = None
    nod: object = None
    total: object = None
    cuadra: bool = True
    ret_opciones: tuple = ()      # valores aceptables si la fila del reporte no cuadra

    def acepta_ret(self, valor):
        return valor == self.ret or valor in self.ret_opciones


@dataclass
class Lectura:
    archivo: str
    entradas: list
    referencias: list
    lineas: list
    fuente: str
    traza: list
    pasada: int = 1
    error: str = ""
    avisos: list = field(default_factory=list)
    paginas: int = 0
    conf_media: float = 1.0
    pag_info: list = field(default_factory=list)   # por pagina: (fuente, renglones, relevante, conf)
    releidas: tuple = ()                            # paginas leidas con esta lectura (las demas: lectura 1)


class OCRNoDisponible(Exception):
    """El PDF es escaneado y el OCR no esta instalado o no carga en este sistema."""


# ------------------------------- lectura -----------------------------------
_ocr = None


_hilos_ocr = None                       # hilos por proceso (None = los que elija onnxruntime)


def load_ocr():
    """rapidocr_onnxruntime (Python <= 3.12) o rapidocr 3.x + onnxruntime. Devuelve
    una funcion imagen -> [(caja, texto, confianza)]."""
    errores = []
    try:
        from rapidocr_onnxruntime import RapidOCR
        motor = RapidOCR(intra_op_num_threads=_hilos_ocr) if _hilos_ocr else RapidOCR()
        return lambda arr: (motor(arr)[0] or [])
    except Exception as e:              # ImportError, o "DLL load failed" en Windows
        errores.append(f"rapidocr_onnxruntime: {e}")
    try:
        import logging
        from rapidocr import RapidOCR
        logging.getLogger("RapidOCR").setLevel(logging.WARNING)
        motor = (RapidOCR(params={"EngineConfig.onnxruntime.intra_op_num_threads": _hilos_ocr})
                 if _hilos_ocr else RapidOCR())

        def leer(arr):
            r = motor(arr)
            return [] if r.boxes is None else list(zip(r.boxes.tolist(), r.txts, r.scores))
        return leer
    except Exception as e:
        errores.append(f"rapidocr: {e}")
    raise OCRNoDisponible("; ".join(errores))


def preprocesar(img, pasada):
    """Imagen para el OCR segun la lectura:
    1: tal cual (300 ppp);
    2: 400 ppp con mas contraste y nitidez;
    3: escaneos malos: imagen ampliada al doble y enfocada (300 ppp x 2).
    No se convierte a blanco y negro: en las pruebas, el OCR (una red neuronal entrenada
    con imagenes en grises) leia mucho peor las paginas binarizadas."""
    if pasada <= 1:
        return img
    from PIL import Image, ImageFilter, ImageOps
    gris = ImageOps.autocontrast(img.convert("L"), cutoff=1)
    if pasada == 2:
        return gris.filter(ImageFilter.SHARPEN).convert("RGB")
    gris = gris.resize((gris.width * 2, gris.height * 2), Image.LANCZOS)
    return gris.filter(ImageFilter.UnsharpMask(radius=2, percent=160, threshold=3)).convert("RGB")


def es_pagina_de_texto(page):
    return len(page.chars) >= 20 or bool(page.chars and not page.images)


def renderizar(page, dpi, lado_max=LADO_MAX_PX):
    """Imagen de la pagina a la resolucion pedida, sin pasar de lado_max pixeles."""
    lado = max(float(page.width), float(page.height)) * dpi / 72.0
    if lado > lado_max:
        dpi = dpi * lado_max / lado
    return page.to_image(resolution=dpi).original


def page_tokens(page, pasada=1):
    """Palabras de la pagina con posicion y confianza; fuente 'texto' u 'OCR'."""
    if es_pagina_de_texto(page):
        words = page.extract_words(x_tolerance=1.5, y_tolerance=2)
        return [Tok(w["x0"], w["x1"], (w["top"] + w["bottom"]) / 2, w["bottom"] - w["top"],
                    w["text"], 1.0, i) for i, w in enumerate(words)], "texto"
    global _ocr
    if _ocr is None:
        _ocr = load_ocr()
    import numpy as np
    try:
        if pasada == 3:
            img = renderizar(page, OCR_DPI, LADO_MAX_PX_AMPLIADA)
        else:
            img = renderizar(page, OCR_DPI_REINTENTO if pasada == 2 else OCR_DPI)
    except ImportError as e:            # sin pypdfium2 (p. ej. BSD) no se puede rasterizar
        raise OCRNoDisponible(str(e)) from e
    img = img.convert("RGB")
    img = preprocesar(img, pasada)
    resultado = list(_ocr(np.array(img)))
    # Inclinacion del escaneo: angulo mediano del borde superior de las cajas de texto.
    import math
    angulos = []
    for box, text, _ in resultado:
        (ax, ay), (bx, by) = box[0], box[1]
        if bx - ax > 3 * abs(box[3][1] - box[0][1]) and len(text) >= 4:
            angulos.append(math.atan2(by - ay, bx - ax))
    ang = sorted(angulos)[len(angulos) // 2] if len(angulos) >= 3 else 0.0
    cos, sen = math.cos(-ang), math.sin(-ang)
    toks = []
    for n, (box, text, score) in enumerate(resultado):
        box = [(x * cos - y * sen, x * sen + y * cos) for x, y in box]   # enderezar
        xs, ys = [p[0] for p in box], [p[1] for p in box]
        x0, x1, h = min(xs), max(xs), max(ys) - min(ys)
        y = (min(ys) + max(ys)) / 2
        ancho = max(len(text), 1)
        for m in re.finditer(r"\S+", text):   # una caja de OCR -> palabras
            toks.append(Tok(x0 + (x1 - x0) * m.start() / ancho, x0 + (x1 - x0) * m.end() / ancho,
                            y, h, m.group(), float(score), n))
    return toks, "OCR"


def agrupar_renglones(toks):
    if not toks:
        return []
    hs = sorted(t.h for t in toks)
    hmed = hs[len(hs) // 2]
    renglones = []
    for t in sorted(toks, key=lambda t: t.y):
        r = renglones[-1] if renglones else None
        if r and abs(t.y - r.y) < 0.6 * max(hmed, min(t.h, r.h)):
            r.toks.append(t)
            ys = sorted(x.y for x in r.toks)
            r.y = ys[len(ys) // 2]
            r.h = sorted(x.h for x in r.toks)[len(r.toks) // 2]
        else:
            r = Renglon(t.y, t.h)
            r.toks.append(t)
            renglones.append(r)
    for r in renglones:
        r.toks.sort(key=lambda t: t.x0)
        r.texto = " ".join(t.text for t in r.toks)
    return renglones


def limpiar_celda(texto):
    s = texto.strip().strip(MARCAS + " ")
    s = re.sub(r"^(\d[\d.,]*)\s+[vVyYrR✓✔√/]$", r"\1", s)     # palomita escrita a mano
    if re.fullmatch(r"\d[\d.,]*[.,]", s):
        s = s.rstrip(".,")                                            # "2,481." -> "2,481"
    return s


def tipo_celda(texto):
    """('lote'|'fecha'|'dinero'|'entero'|'cajas'|'codigo'|'texto'|'vacio', valor)."""
    s = limpiar_celda(texto)
    if not s:
        return "vacio", None
    lotes = lotes_en(s)
    if lotes and re.fullmatch(r"[A-Za-z]{0,4}[\s:#.\-]*" + re.escape(lotes[0]) + r"[\s.]*", s,
                              re.I if LOTE_MAYUSCULAS else 0):
        return "lote", lotes[0]
    if FECHA_RE.match(s):
        return "fecha", s
    if DINERO_RE.match(s):
        return "dinero", a_dinero(s)
    v = a_entero(s)
    if v is not None:
        return "entero", v
    m = UNIDAD_RE.match(s)
    if m and a_entero(m.group(1)) is not None:
        return ("cajas" if m.group(2).lower().startswith(CAJA_PREFIJOS) else "entero"), a_entero(m.group(1))
    if re.search(r"[A-Za-z]", s) and re.search(r"\d", s) and " " not in s:
        return "codigo", s
    return "texto", s


def es_valor(texto):
    return tipo_celda(texto)[0] in ("entero", "dinero", "fecha", "lote", "cajas")


def es_marca(texto):
    return bool(re.fullmatch(r"[✓✔√/\\|*vVyY]", texto.strip()))


def unir_lotes_partidos(toks):
    """El OCR a veces parte un lote escrito a mano: "262402 54" -> "26240254"."""
    res = []
    for t in toks:
        if res:
            p = res[-1]
            junto = p.text + t.text
            if (p.text.isdigit() and t.text.isdigit() and lotes_en(junto) == [junto]
                    and not lotes_en(p.text) and t.x0 - p.x1 < 1.2 * max(p.h, t.h)):
                res[-1] = Tok(p.x0, t.x1, p.y, max(p.h, t.h), junto, min(p.conf, t.conf), p.caja)
                continue
        res.append(t)
    return res


def armar_celdas(renglon):
    """Une palabras cercanas en celdas, sin juntar nunca dos valores numericos."""
    grupos = []
    for t in unir_lotes_partidos(renglon.toks):
        if grupos:
            g = grupos[-1]
            ult = g[-1]
            gtexto = " ".join(x.text for x in g)
            hueco = t.x0 - ult.x1
            h = max(t.h, ult.h)
            misma_caja = t.caja == ult.caja and t.conf < 1.0
            if gtexto.strip() == "$" and re.match(r"\d", t.text) and hueco < 12 * h:
                g.append(t)
                continue
            if es_marca(t.text) and hueco < 1.5 * h and re.search(r"\d", gtexto):
                g.append(t)
                continue
            juntar = not (es_valor(gtexto) and es_valor(t.text)) and not lotes_en(t.text) \
                and not lotes_en(gtexto) and (misma_caja or hueco < 0.5 * h)
            if juntar:
                g.append(t)
                continue
        grupos.append([t])
    celdas = []
    for g in grupos:
        texto = " ".join(x.text for x in g)
        tipo, valor = tipo_celda(texto)
        celdas.append(Celda(g[0].x0, g[-1].x1, texto, min(x.conf for x in g), tipo, valor))
    return celdas


def solape(a0, a1, b0, b1):
    return min(a1, b1) - max(a0, b0)


# ----------------------- formularios de lado -------------------------------
def analizar_de_lado(renglones, celdas, pnum, tipo_pag, usados, traza):
    """Un renglon 'Lote' con un lote por columna y renglones de cantidad alineados."""
    entradas = []
    for i, r in enumerate(renglones):
        cs = celdas[i]
        etiqueta = " ".join(c.text for c in cs if c.tipo in ("texto", "codigo") and c.x1 <= min(
            (x.x0 for x in cs if x.tipo == "lote"), default=0) + 1)
        cols = [((c.x0 + c.x1) / 2, c.valor) for c in cs if c.tipo == "lote"]
        if len(cols) < 2 or not LOTE_PALABRA.search(etiqueta):
            continue
        xs = sorted(c[0] for c in cols)
        tol = min(b - a for a, b in zip(xs, xs[1:])) * 0.5
        mejor = None
        for j in range(max(0, i - 15), min(len(renglones), i + 25)):
            if j == i:
                continue
            cj = celdas[j]
            primero = min((c.x0 for c in cj if c.tipo == "entero"), default=None)
            if primero is None:
                continue
            etq = " ".join(c.text for c in cj if c.x1 <= primero and c.tipo in ("texto", "codigo"))
            sem = semantica(etq)
            if not (sem["cant"] or sem["ret"] or sem["nod"]):
                continue
            prio = next((k for k, w in enumerate(PRIORIDAD_RENGLON) if w in sin_acentos(etq)),
                        len(PRIORIDAD_RENGLON))
            pares = []
            for c in cj:
                if c.tipo != "entero":
                    continue
                cx = (c.x0 + c.x1) / 2
                col = min(cols, key=lambda k: abs(k[0] - cx))
                if abs(col[0] - cx) <= tol:
                    pares.append((col[1], c.valor, c.conf))
            if pares and (mejor is None or prio < mejor[0]):
                mejor = (prio, pares, etq, j, sem)
        traza.append(f"pagina {pnum}: formulario de lado, renglon '{etiqueta}' con lotes "
                     f"{[c[1] for c in cols]}" + (f"; cantidades del renglon '{mejor[2]}'" if mejor
                                                   else "; ningun renglon de cantidad con numeros"))
        usados.add(i)
        if mejor:
            usados.add(mejor[3])
            tipo = "nod" if mejor[4]["nod"] else ("ret" if mejor[4]["ret"] else tipo_pag)
            for lote, v, conf in mejor[1]:
                entradas.append(Entrada("", pnum, lote, v, tipo, "de lado", conf,
                                        notas=[f"renglon '{mejor[2]}'"]))
    return entradas


# -------------------------------- tablas -----------------------------------
def filas_con_items(renglones, celdas, usados):
    """Renglones con un lote y al menos un entero. Tambien acepta el lote en el renglon
    siguiente, debajo de la descripcion ("TORNILLO ... / LOTE: 25240310")."""
    items, saltar = [], set()
    for i, cs in enumerate(celdas):
        if i in usados or i in saltar:
            continue
        lotes = [c for c in cs if c.tipo == "lote"] + [c for c in cs if c.tipo != "lote" and lotes_en(c.text)]
        enteros = [c for c in cs if c.tipo == "entero"]
        if lotes and enteros:
            lc = lotes[0]
            items.append(Item(i, lc.valor if lc.tipo == "lote" else lotes_en(lc.text)[0], lc, cs))
        elif enteros and not lotes and len(cs) >= 3:
            for j in (i + 1, i + 2):
                if j >= len(celdas) or j in usados:
                    break
                cj = celdas[j]
                if any(c.tipo == "entero" for c in cj):
                    break
                lc = next((c for c in cj if lotes_en(c.text)), None)
                if lc and len(cj) <= 2 and any(c.tipo in ("texto", "codigo") and
                                               solape(c.x0, c.x1, lc.x0, lc.x1) > 0 for c in cs):
                    items.append(Item(i, lotes_en(lc.text)[0], lc, cs))
                    saltar.add(j)
                    break
    return items, saltar


def agrupar_tablas(items, renglones):
    if not items:
        return []
    hmed = sorted(r.h for r in renglones)[len(renglones) // 2]
    tablas, actual = [], [items[0]]
    for it in items[1:]:
        if renglones[it.fila].y - renglones[actual[-1].fila].y > 5 * hmed:
            tablas.append(actual)
            actual = []
        actual.append(it)
    tablas.append(actual)
    return tablas


def columnas_de(tabla):
    """Agrupa las celdas de los renglones de la tabla en columnas por posicion."""
    cols = []
    for k, it in enumerate(tabla):
        for c in it.celdas:
            w = max(c.x1 - c.x0, 1e-6)
            mejor, ov_mejor = None, 0
            for col in cols:
                ov = solape(c.x0, c.x1, col["x0"], col["x1"])
                if ov > ov_mejor:
                    mejor, ov_mejor = col, ov
            if mejor is not None and ov_mejor >= 0.4 * min(w, mejor["x1"] - mejor["x0"]):
                mejor["x0"], mejor["x1"] = min(mejor["x0"], c.x0), max(mejor["x1"], c.x1)
                mejor["celdas"].setdefault(k, []).append(c)
            else:
                cols.append({"x0": c.x0, "x1": c.x1, "celdas": {k: [c]}})
    cols.sort(key=lambda c: c["x0"])
    for col in cols:                    # varias celdas del mismo renglon -> una
        for k, cs in list(col["celdas"].items()):
            if len(cs) > 1:
                texto = " ".join(c.text for c in sorted(cs, key=lambda c: c.x0))
                tipo, valor = tipo_celda(texto)
                col["celdas"][k] = [Celda(min(c.x0 for c in cs), max(c.x1 for c in cs), texto,
                                          min(c.conf for c in cs), tipo, valor)]
        col["celda"] = {k: cs[0] for k, cs in col["celdas"].items()}
    return cols


def encabezados(cols, tabla, renglones, celdas):
    """Texto de encabezado de cada columna: renglones sin numeros justo encima de la
    tabla (hasta 3, para encabezados en dos lineas) y la etiqueta a la izquierda."""
    primero = tabla[0].fila
    hmed = renglones[primero].h
    zona, y_prev = [], renglones[primero].y
    for j in range(primero - 1, max(-1, primero - 4), -1):
        r = renglones[j]
        if lotes_en(r.texto) or y_prev - r.y > 3.5 * max(hmed, r.h):
            break
        if sum(1 for c in celdas[j] if c.tipo in ("entero", "dinero")) > 1:
            break
        zona.append(r)
        y_prev = r.y
    textos = [[] for _ in cols]
    for r in zona:
        for t in r.toks:
            cx = (t.x0 + t.x1) / 2
            dentro = [k for k, c in enumerate(cols) if c["x0"] - 0.5 * t.h <= cx <= c["x1"] + 0.5 * t.h]
            if dentro:
                k = min(dentro, key=lambda k: abs((cols[k]["x0"] + cols[k]["x1"]) / 2 - cx))
            else:
                k = min(range(len(cols)), key=lambda k: abs((cols[k]["x0"] + cols[k]["x1"]) / 2 - cx))
                if abs((cols[k]["x0"] + cols[k]["x1"]) / 2 - cx) > 3 * t.h + (cols[k]["x1"] - cols[k]["x0"]):
                    continue
            textos[k].append(t.text)
    for k, col in enumerate(cols):      # etiqueta inmediatamente a la izquierda ("Piezas: 30")
        for i, c in col["celda"].items():
            fila = tabla[i].celdas
            izq = [x for x in fila if x.x1 <= c.x0 + 1 and x.tipo in ("texto", "codigo")]
            if izq:
                etq = max(izq, key=lambda x: x.x1)
                if etq.text.rstrip().endswith(":") or semantica(etq.text)["cant"]:
                    textos[k].append(etq.text)
    return [" ".join(t) for t in textos]


def es_secuencia(vals):
    return len(vals) >= 2 and all(b - a == 1 for a, b in zip(vals, vals[1:]))


def relaciones_suma(M, cand):
    """Ternas (a, b, s) con v[a] + v[b] == v[s] en todos los renglones (o en >= 80%)."""
    res = []
    col = {k: [fila[k] for fila in M] for k in cand}
    trivial = {k: all(v in (0, None) for v in col[k]) for k in cand}
    pares = [(a, b) for a in cand for b in cand if a < b and not trivial[a] and not trivial[b]]
    for s in cand:
        vs = col[s]
        for a, b in pares:
            if s == a or s == b:
                continue
            ok = completos = 0
            for x, y, z in zip(col[a], col[b], vs):
                if x is not None and y is not None and z is not None:
                    completos += 1
                    if x + y == z:
                        ok += 1
            if completos >= 2 and ok >= max(2, completos - (completos // 5)):
                res.append((a, b, s))
    return res


def roles_de_reporte(M, cand, sems, sumas):
    """Reporte de retiro: fabricado = distribuido + no distribuido, y
    total = no distribuido + retirado. El sumando comun es lo no distribuido; el otro
    sumando menor es lo retirado (no se puede retirar mas de lo distribuido)."""
    for i, (a1, b1, s1) in enumerate(sumas):
        for a2, b2, s2 in sumas[i + 1:]:
            comun = {a1, b1} & {a2, b2}
            if len(comun) != 1 or s1 == s2:
                continue
            x = comun.pop()
            p = b1 if a1 == x else a1
            q = b2 if a2 == x else a2
            if p == q:
                continue
            filas = [f for f in M if f[p] is not None and f[q] is not None]
            if filas and all(f[p] <= f[q] for f in filas):
                return {"nod": x, "ret": p, "total": s1, "dist": q, "fab": s2}
            if filas and all(f[q] <= f[p] for f in filas):
                return {"nod": x, "ret": q, "total": s2, "dist": p, "fab": s1}
    for a, b, s in sumas:               # una sola suma: hace falta el encabezado
        if sems[s]["total"]:
            if sems[a]["nod"] or sems[b]["ret"]:
                return {"nod": a, "ret": b, "total": s}
            if sems[b]["nod"] or sems[a]["ret"]:
                return {"nod": b, "ret": a, "total": s}
    return None


def fila_de_reporte(f, roles, pnum, lote, traza):
    """Una fila del reporte. Si un valor falta o esta mal leido, las dos relaciones
    (fabricado - distribuido = no distribuido; no distribuido + retirado = total)
    permiten recuperarlo."""
    g = lambda k: f[roles[k]] if k in roles else None
    fab, dist, nod, ret, tot = g("fab"), g("dist"), g("nod"), g("ret"), g("total")
    opciones_nod = [x for x in (nod, None if None in (fab, dist) else fab - dist,
                                None if None in (tot, ret) else tot - ret) if x is not None and x >= 0]
    nod_ok = Counter(opciones_nod).most_common(1)[0][0] if opciones_nod else None
    cuadra = None not in (nod_ok, ret, tot) and nod_ok + ret == tot and nod_ok == nod
    opciones = ()
    if not cuadra:
        if ret is None and None not in (tot, nod_ok):
            ret = tot - nod_ok
        elif None not in (tot, nod_ok, ret):
            opciones = (tot - nod_ok,)
        traza.append(f"pagina {pnum}: fila del reporte del lote {lote} no cuadra "
                     f"(no distribuido {nod}, retirado {ret}, total {tot}); se acepta tambien "
                     f"retirado = {fmt(tot - nod_ok) if None not in (tot, nod_ok) else '-'}")
    return Referencia("", pnum, lote, ret, nod_ok, tot, bool(cuadra), opciones)


def analizar_tablas(renglones, celdas, pnum, tipo_pag, usados, traza):
    entradas, referencias = [], []
    items, saltados = filas_con_items(renglones, celdas, usados)
    tablas = agrupar_tablas(items, renglones)
    filas_items = {it.fila for it in items} | saltados
    # Numeros fuera de las tablas: posibles totales impresos del documento.
    enteros_doc, dinero_doc, totales_rotulados = [], [], []
    for i, cs in enumerate(celdas):
        if i in filas_items or any(c.tipo == "fecha" for c in cs):
            continue                      # renglones de datos (con fecha) no son totales
        ents = [c.valor for c in cs if c.tipo == "entero"]
        enteros_doc += ents
        dinero_doc += [c.valor for c in cs if c.tipo == "dinero"]
        t = sin_acentos(renglones[i].texto)
        if ents and TOTAL_PALABRA_RE.search(t) and FILA_TOTAL_RE.search(t):
            totales_rotulados.append(max(ents))

    for tabla in tablas:
        cols = columnas_de(tabla)
        heads = encabezados(cols, tabla, renglones, celdas)
        sems = [semantica(h) for h in heads]
        n = len(tabla)
        col_lote = next((k for k, c in enumerate(cols)
                         if any(x is it.celda_lote for i, it in enumerate(tabla) for x in [c["celda"].get(i)])), None)
        # Columnas de enteros candidatas a cantidad
        cand = []
        for k, col in enumerate(cols):
            if k == col_lote:
                continue
            tipos = Counter(c.tipo for c in col["celda"].values() if c.tipo != "vacio")
            if not tipos or tipos["entero"] < 0.6 * sum(tipos.values()) or len(col["celda"]) < 0.5 * n:
                continue
            vals = [col["celda"][i].valor if i in col["celda"] and col["celda"][i].tipo == "entero" else None
                    for i in range(n)]
            conocidos = [v for v in vals if v is not None]
            if n >= 2 and es_secuencia(conocidos):
                continue                  # numero de renglon
            if n >= 2 and len(set(conocidos)) == 1 and conocidos[0] >= 10000:
                continue                  # codigo repetido (p. ej. CODIGO SAP)
            if all(1990 <= v <= 2100 for v in conocidos) and not sems[k]["cant"]:
                continue                  # anios
            cand.append(k)
        M = [[(cols[k]["celda"][i].valor if i in cols[k]["celda"] and cols[k]["celda"][i].tipo == "entero"
               else None) if k in cand else None for k in range(len(cols))] for i in range(n)]

        # 1) Reporte de retiro (relaciones aritmeticas entre columnas)
        roles = None
        if len(cand) >= 3 and n >= 2:
            sumas = relaciones_suma(M, cand)
            if sumas:
                roles = roles_de_reporte(M, cand, sems, sumas)
        if roles:
            traza.append(f"pagina {pnum}: reporte de retiro; columnas: no distribuido '{heads[roles['nod']]}', "
                         f"retirado '{heads[roles['ret']]}', total '{heads[roles['total']]}'")
            for i, it in enumerate(tabla):
                referencias.append(fila_de_reporte(M[i], roles, pnum, it.lote, traza))
            continue

        # 2) Tabla normal: elegir la columna de cantidad
        puntos = {}
        detalles = {}
        for k in cand:
            p, por = 0.0, []
            s = sems[k]
            if s["cant"] or s["ret"] or s["nod"]:
                p += 2
                por.append("encabezado de cantidad")
            if s["ret"] or s["nod"]:
                p += 1
            if s["orig"] and not (s["ret"] or s["nod"]):
                p -= 3
                por.append("cantidad original (facturada/distribuida)")
            if s["caja"] and not PIEZA_RE.search(sin_acentos(heads[k])):
                p -= 4
                por.append("cuenta cajas")
            if s["nocant"] and not s["cant"]:
                p -= 3
                por.append("encabezado de motivo, folio o codigo")
            vals = [M[i][k] for i in range(n)]
            rel = 0
            for i, it in enumerate(tabla):
                din = sorted(c.valor for c in it.celdas if c.tipo == "dinero")
                if vals[i] is not None and any(abs(vals[i] * a - b) <= 0.011 for a in din for b in din if a < b):
                    rel += 1
            if rel:
                p += 4
                por.append(f"piezas x precio = importe en {rel} renglon(es)")
            suma = sum(v for v in vals if v is not None)
            if suma >= 10 and suma in enteros_doc:
                p += 3
                por.append(f"suma {fmt(suma)} = numero impreso en el documento")
            if len(cand) == 1 and n >= 2:
                p += 1
                por.append("unica columna de enteros")
            if all(v is not None and v >= 1_000_000 for v in vals):
                p -= 5
            conocidos = [v for v in vals if v is not None]
            if conocidos and max(conocidos) > max(MAX_RETIRADO_POR_LOTE, MAX_NO_DISTRIBUIDA_POR_LOTE):
                p -= 4
                por.append("valores mayores que el tope por lote")
            if n >= 3 and len(set(conocidos)) == 1:
                p -= 2
                por.append("mismo valor en todos los renglones")
            puntos[k], detalles[k] = p, por
        if not puntos:
            traza.append(f"pagina {pnum}: tabla de {n} renglon(es) sin columna de enteros utilizable")
            continue
        k = max(puntos, key=lambda c: (puntos[c], -cols[c]["x0"]))
        minimo = 2 if n == 1 else 1
        if puntos[k] < minimo:
            traza.append(f"pagina {pnum}: tabla de {n} renglon(es) descartada: ninguna columna de cantidad "
                         f"clara ({', '.join(repr(heads[c]) for c in cand) or 'sin columnas de enteros'})")
            continue
        empate = [c for c in puntos if c != k and puntos[c] == puntos[k]]
        s = sems[k]
        tipo = "nod" if (s["nod"] and not s["ret"]) else ("ret" if s["ret"] else tipo_pag)
        traza.append(f"pagina {pnum}: tabla de {n} renglon(es); columna de cantidad '{heads[k]}' "
                     f"({'; '.join(detalles[k]) or 'por estructura'})"
                     + (f"; EMPATE con {[heads[c] for c in empate]}" if empate else ""))
        nuevas = []
        for i, it in enumerate(tabla):
            c = cols[k]["celda"].get(i)
            if c is None or c.tipo != "entero":
                continue
            e = Entrada("", pnum, it.lote, c.valor, tipo, "tabla", c.conf)
            din = sorted(x.valor for x in it.celdas if x.tipo == "dinero")
            par = next(((a, b) for a in din for b in din if a < b and abs(c.valor * a - b) <= 0.011), None)
            if par:
                e.verificado.append(f"piezas x precio = importe ({fmt(c.valor)} x {par[0]:,.2f} = {par[1]:,.2f})")
            e.precio = par[0] if par else (din[0] if len(din) == 1 else None)
            e.dinero = din
            if empate:
                e.notas.append("columna de cantidad dudosa (empate)")
            nuevas.append(e)
        verificar_tabla(nuevas, enteros_doc, dinero_doc, totales_rotulados)
        entradas += nuevas
        usados.update(it.fila for it in tabla)
    usados.update(saltados)
    return entradas, referencias


def sin_marca(v):
    """Valores posibles si una raya o palomita a mano se leyo como 1, 7 o 4."""
    t = str(v)
    res = set()
    if len(t) > 1 and t[0] in "17":
        res.add(int(t[1:]))
    if len(t) > 1 and t[-1] in "174":
        res.add(int(t[:-1]))
    return [x for x in res if x > 0]


def verificar_tabla(nuevas, enteros_doc, dinero_doc, totales_rotulados):
    """Compara la suma de la tabla con los totales impresos. Si el total rotulado de
    piezas no coincide y solo un renglon no esta verificado, lo corrige cuando el
    precio confirma el valor corregido."""
    if not nuevas:
        return
    suma = sum(e.piezas for e in nuevas)
    if suma >= 10 and suma in enteros_doc:
        for e in nuevas:
            e.verificado.append(f"suma de la tabla = total impreso ({fmt(suma)})")
    precios = [getattr(e, "precio", None) for e in nuevas]
    if all(p is not None for p in precios):
        importe = round(sum(e.piezas * p for e, p in zip(nuevas, precios)), 2)
        if any(abs(importe - d) <= 0.011 for d in dinero_doc):
            for e in nuevas:
                e.verificado.append(f"suma de piezas x precio = importe total impreso ({importe:,.2f})")
    for total in totales_rotulados:
        if total == suma:
            continue
        sin_verif = [e for e in nuevas if not e.verificado]
        if len(sin_verif) == 1:
            e = sin_verif[0]
            correcto = total - (suma - e.piezas)
            p = getattr(e, "precio", None)
            din = getattr(e, "dinero", [])
            con_renglon = any(abs(correcto * a - b) <= 0.011 for a in din for b in din if a < b)
            con_total = bool(p) and any(abs(correcto * p - d) <= 0.011 for d in dinero_doc)
            if correcto > 0 and (con_renglon or con_total):
                e.notas.append(f"CORREGIDO: se leyo {fmt(e.piezas)}; el total impreso ({fmt(total)}) "
                               f"y el importe confirman {fmt(correcto)}")
                e.piezas = correcto
                e.verificado.append("corregido con el total y el importe impresos")
                continue
        cand = [(e, v) for e in sin_verif for v in sin_marca(e.piezas) if suma - e.piezas + v == total]
        if len(cand) == 1:
            e, v = cand[0]
            e.notas.append(f"CORREGIDO: se leyo {fmt(e.piezas)}; una marca a mano junto al numero se leyo como "
                           f"digito; con {fmt(v)} la tabla suma el total impreso ({fmt(total)})")
            e.piezas = v
            e.verificado.append("corregido con el total de piezas impreso")
            continue
        for e in nuevas:
            e.fallas.append(f"la suma de la tabla ({fmt(suma)}) no coincide con el total de piezas "
                            f"impreso ({fmt(total)})")


# --------------------------- lineas con etiquetas --------------------------
def analizar_etiquetas(renglones, pnum, tipo_pag, usados, traza):
    ev = []
    for i, r in enumerate(renglones):
        if i in usados:
            continue
        e = [(m.start(), "L", lotes_en(m.group("id"))[0] if lotes_en(m.group("id")) else None, 1.0)
             for m in LOTE_ETIQUETA.finditer(r.texto)]
        e += [(m.start(), "C", a_entero(m.group("n").rstrip(".,")), min(t.conf for t in r.toks))
              for m in CANT_ETIQUETA.finditer(r.texto)]
        ev += [(k, v, c) for _, k, v, c in sorted(e, key=lambda x: x[0]) if v is not None]
    entradas = []
    if not ev:
        return entradas
    primero, pend = ev[0][0], None
    for kind, val, conf in ev:
        if kind == primero:
            pend = (val, conf)
        elif pend is not None:
            lote, v, c = (pend[0], val, conf) if primero == "L" else (val, pend[0], pend[1])
            entradas.append(Entrada("", pnum, lote, v, tipo_pag, "etiquetas", c))
            pend = None
    if entradas:
        traza.append(f"pagina {pnum}: lineas con etiquetas: "
                     + ", ".join(f"lote {e.lote} = {fmt(e.piezas)}" for e in entradas))
    return entradas


# ------------------------------- un PDF ------------------------------------
def analizar_pagina(renglones, pnum, traza):
    celdas = [armar_celdas(r) for r in renglones]
    tipo_pag = tipo_documento(" ".join(r.texto for r in renglones))
    usados = set()
    entradas = analizar_de_lado(renglones, celdas, pnum, tipo_pag, usados, traza)
    e2, refs = analizar_tablas(renglones, celdas, pnum, tipo_pag, usados, traza)
    entradas += e2
    entradas += analizar_etiquetas(renglones, pnum, tipo_pag, usados, traza)
    return entradas, refs, tipo_pag


def deduplicar_en_archivo(entradas, avisos, traza):
    """El mismo lote repetido en el mismo PDF no se suma dos veces.

    1) Dos paginas con exactamente los mismos lotes son el mismo documento (p. ej. la
       devolucion por lote y su detalle, o el acuse impreso dos veces): se cuenta una,
       de preferencia la que cuadra con su total impreso.
    2) Un lote en varias paginas con la misma cantidad se cuenta una vez; una pagina que
       es la suma de las otras no se cuenta."""
    paginas = defaultdict(list)
    for e in entradas:
        paginas[e.pagina].append(e)
    nums = sorted(paginas)
    for i, a in enumerate(nums):
        for b in nums[i + 1:]:
            ea = [e for e in paginas[a] if not e.excluido]
            eb = [e for e in paginas[b] if not e.excluido]
            la, lb = {(e.tipo, e.lote) for e in ea}, {(e.tipo, e.lote) for e in eb}
            if len(la) < 2 or la != lb:
                continue
            ver = lambda es: sum(1 for e in es if any("total impreso" in v for v in e.verificado))
            sa, sb = sum(e.piezas for e in ea), sum(e.piezas for e in eb)
            quedar, quitar = (a, eb) if ver(ea) >= ver(eb) else (b, ea)
            for e in quitar:
                e.excluido = f"la pagina {e.pagina} repite los mismos lotes que la pagina {quedar}"
            traza.append(f"paginas {a} y {b}: mismos lotes ({len(la)}); se cuenta la pagina {quedar}")
            if sa != sb:
                avisos.append(f"las paginas {a} y {b} tienen los mismos lotes pero suman distinto "
                              f"({fmt(sa)} y {fmt(sb)}); se conto la pagina {quedar}"
                              + (" porque cuadra con su total impreso" if ver(ea) != ver(eb) else ""))
    por = defaultdict(list)
    for e in entradas:
        por[(e.tipo, e.lote)].append(e)
    for (tipo, lote), es in por.items():
        if len(es) < 2:
            continue
        paginas = defaultdict(list)
        for e in es:
            paginas[e.pagina].append(e)
        for p, pes in paginas.items():          # dentro de una pagina: renglon de total
            activos = [e for e in pes if not e.excluido]
            if len(activos) >= 3:
                for e in activos:
                    if e.piezas == sum(x.piezas for x in activos if x is not e):
                        e.excluido = "renglon de total del lote"
                        break
        sumas = {p: sum(e.piezas for e in pes if not e.excluido) for p, pes in paginas.items()
                 if any(not e.excluido for e in pes)}
        if len(sumas) < 2:
            continue
        orden = sorted(sumas)
        valores = [sumas[p] for p in orden]
        if len(set(valores)) == 1 or DUPLICADOS_EN_ARCHIVO == "todos":
            if len(set(valores)) == 1:
                quedar = orden[0]
            else:
                quedar = max(orden, key=lambda p: (sumas[p], -p))
            for p in orden:
                if p != quedar:
                    for e in paginas[p]:
                        if not e.excluido:
                            e.excluido = f"lote repetido en la pagina {quedar} del mismo archivo"
            traza.append(f"lote {lote}: aparece en las paginas {orden}; se cuenta una vez (pagina {quedar})")
            continue
        total_pag = next((p for p in orden if sumas[p] == sum(sumas[q] for q in orden if q != p)), None)
        if total_pag is not None:
            for e in paginas[total_pag]:
                if not e.excluido:
                    e.excluido = "pagina con el total del lote (suma de las otras paginas)"
            traza.append(f"lote {lote}: la pagina {total_pag} es la suma de las otras; no se cuenta")
            continue
        avisos.append(f"lote {lote} aparece en varias paginas con cantidades distintas "
                      f"({', '.join(f'p{p}: {fmt(sumas[p])}' for p in orden)}); se suman")


RELEVANTE_RE = re.compile(r"\d{6,}|l[o0]te|pieza|pzas|cantidad|devol|retir", re.I)
CASI_LOTE_RE = re.compile(r"(?<!\d)2[0-9]{6,8}(?!\d)")   # numeros que podrian ser lotes mal leidos


def leer_pdf(ruta, nombre, pasada=1, traza=None):
    """Lee y analiza un PDF completo en este proceso (sin cache). lotes_auto.py usa
    LectorPaginas (paralelo y con cache) y analizar_documento; diagnostico.py usa esto."""
    with pdfplumber.open(ruta) as doc:
        paginas = [page_tokens(page, pasada) for page in doc.pages]
    return analizar_documento(nombre, paginas, pasada, traza)


def analizar_documento(nombre, paginas, pasada=1, traza=None, releidas=()):
    """paginas: [(tokens, fuente)] en orden. Devuelve la Lectura del PDF."""
    traza = [] if traza is None else traza
    entradas, refs, lineas, fuentes, tipos, avisos, confs = [], [], [], set(), [], [], []
    npag = len(paginas)
    info = []
    if releidas:
        traza.append(f"lectura {pasada} de las paginas {', '.join(map(str, releidas))}"
                     + ("; las demas, de la lectura 1" if len(releidas) < npag else ""))
    for pnum, (toks, fuente) in enumerate(paginas, 1):
        fuentes.add(fuente)
        confs += [t.conf for t in toks]
        renglones = agrupar_renglones(toks)
        textos = [r.texto for r in renglones]
        lineas += textos
        traza.append(f"pagina {pnum}: {len(renglones)} renglones ({fuente})")
        e, r, tipo_pag = analizar_pagina(renglones, pnum, traza)
        tipos.append(tipo_pag)
        entradas += e
        refs += r
        cp = [t.conf for t in toks]
        info.append((fuente, len(renglones), bool(RELEVANTE_RE.search(" ".join(textos))),
                     sum(cp) / len(cp) if cp else 0.0))
    # paginas sin tipo reconocido: tipo mayoritario del archivo
    conocido = Counter(t for t in tipos if t)
    tipo_archivo = conocido.most_common(1)[0][0] if conocido else None
    for e in entradas:
        e.archivo = nombre
        if e.tipo is None:
            e.tipo = tipo_archivo or "ret"
            if tipo_archivo is None:
                e.notas.append("tipo de documento no reconocido: se cuenta como retirado")
    for r in refs:
        r.archivo = nombre
    deduplicar_en_archivo(entradas, avisos, traza)
    fuente = "mixto" if len(fuentes) > 1 else (fuentes.pop() if fuentes else "")
    return Lectura(nombre, entradas, refs, lineas, fuente, traza, pasada, avisos=avisos, paginas=npag,
                   conf_media=sum(confs) / len(confs) if confs else 0.0, pag_info=info,
                   releidas=tuple(releidas))


def paginas_dudosas(L, problema_lotes):
    """Paginas escaneadas de un PDF dudoso que vale la pena leer otra vez: las que tienen
    cantidades sin verificar, con fallas, de lotes que no cuadran o con lotes fuera del
    Excel, filas de reporte que no cuadran, y las que no dieron cantidades pero tienen
    texto de lotes o piezas, se leyeron con poca confianza o casi vacias. Si el PDF no
    dio nada, todas sus paginas escaneadas. Las paginas cuyas cantidades estan todas
    verificadas no se vuelven a leer: otra lectura no puede mejorarlas."""
    ocr = [p for p, (fuente, *_ ) in enumerate(L.pag_info, 1) if fuente == "OCR"]
    activas = [e for e in L.entradas if not e.excluido]
    if not activas and not L.referencias:
        return ocr
    sel = set()
    con_datos = {e.pagina for e in L.entradas} | {r.pagina for r in L.referencias}
    for e in L.entradas:
        if (e.excluido == FUERA_EXCEL or e.lote in problema_lotes or e.fallas
                or (not e.excluido and not e.verificado)
                or (not e.excluido and e.conf < CONFIANZA_MINIMA and not e.verificado)):
            sel.add(e.pagina)
    sel |= {r.pagina for r in L.referencias if not r.cuadra}
    for p in ocr:
        if p in con_datos:
            continue
        _, nren, relevante, conf = L.pag_info[p - 1]
        if relevante or conf < CONFIANZA_MINIMA or nren < 5:
            sel.add(p)
    return sorted(p for p in sel if p in ocr)


def calidad(L):
    """Orden para elegir entre lecturas: mas cantidades verificadas, menos fallas, mas
    cantidades leidas, mas confianza del OCR."""
    ents = [e for e in L.entradas if not e.excluido]
    confs = [e.conf for e in ents] or [0]
    return (sum(1 for e in ents if e.verificado), -sum(len(e.fallas) for e in ents),
            len(ents) + len(L.referencias), sum(confs) / len(confs))


def dudosa(L):
    ents = [e for e in L.entradas if not e.excluido]
    if L.fuente in ("OCR", "mixto") and not ents and not L.referencias:
        return True
    return any(e.fallas or (not e.verificado and e.conf < CONFIANZA_MINIMA) for e in ents)


# --------------------------- Excel de referencia ---------------------------
@dataclass
class ExcelRef:
    ruta: str
    hoja: str
    col_lote: str
    col_nod: str
    col_ret: str
    refs: dict
    avisos: list


def leer_excel_referencia(ruta):
    """Busca en el libro una hoja con una columna 'no distribuida' y lee, por lote, la
    cantidad no distribuida y (si existe) la cantidad retirada segun clientes. Devuelve
    None si el libro no tiene esa columna."""
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    global LOTE_ID
    wb = load_workbook(ruta, data_only=True, read_only=True)
    modo_lista = LOTE_ID is None        # perfil sin formato de lote: los lotes salen de este Excel
    for ws in wb.worksheets:
        filas = list(ws.iter_rows(values_only=True))
        c_nod = c_ret = fila_enc = c_lista = None
        for i, fila in enumerate(filas[:60]):
            for j, v in enumerate(fila):
                t = sin_acentos(str(v)) if v is not None else ""
                if c_nod is None and EXCEL_NOD_RE.search(t):
                    c_nod, fila_enc = j, i
                elif c_ret is None and EXCEL_RET_RE.search(t) and not EXCEL_RET_EXCLUYE.search(t):
                    c_ret = j
                    fila_enc = i if fila_enc is None else max(fila_enc, i)
                elif modo_lista and c_lista is None and EXCEL_LOTE_RE is not None and EXCEL_LOTE_RE.search(t):
                    c_lista = j
                    fila_enc = i if fila_enc is None else max(fila_enc, i)
        if c_nod is None and not (modo_lista and c_lista is not None and c_ret is not None):
            continue
        if modo_lista:
            if c_lista is None:
                continue
            valores = set()
            for fila in filas[fila_enc + 1:]:
                v = fila[c_lista] if c_lista < len(fila) else None
                v = str(int(v)) if isinstance(v, float) and v.is_integer() else ("" if v is None else str(v))
                v = v.strip().upper()
                if v and len(v) <= 40 and re.search(r"\d", v):
                    valores.add(v)
            if not valores:
                continue
            poner_patron_lote("|".join(re.escape(x) for x in sorted(valores, key=lambda x: (-len(x), x))))
        # columna de lote: la que tiene mas lotes validos debajo del encabezado
        cuenta = Counter()
        for fila in filas[fila_enc + 1:]:
            for j, v in enumerate(fila):
                if v is not None and lotes_en(str(int(v)) if isinstance(v, float) and v.is_integer() else str(v)):
                    cuenta[j] += 1
        if not cuenta:
            continue
        c_lote = cuenta.most_common(1)[0][0]
        refs, avisos = {}, []
        sin_valor = 0
        for n, fila in enumerate(filas[fila_enc + 1:], fila_enc + 2):
            v = fila[c_lote] if c_lote < len(fila) else None
            if v is None:
                continue
            lotes = lotes_en(str(int(v)) if isinstance(v, float) and v.is_integer() else str(v))
            if not lotes:
                continue
            def num(c):
                x = fila[c] if c is not None and c < len(fila) else None
                if isinstance(x, (int, float)):
                    return int(round(x))
                if isinstance(x, str) and a_entero(x.strip()) is not None:
                    return a_entero(x.strip())
                return None
            nod, ret = num(c_nod), num(c_ret)
            if nod is None and c_nod is not None and c_nod < len(fila) and fila[c_nod] not in (None, ""):
                sin_valor += 1
            lote = lotes[0]
            if lote in refs:
                avisos.append(f"el lote {lote} aparece dos veces en el Excel (fila {n}); se suman")
                r = refs[lote]
                r.nod = (r.nod or 0) + (nod or 0)
                r.ret = None if r.ret is None and ret is None else (r.ret or 0) + (ret or 0)
                continue
            refs[lote] = Referencia(pathlib.Path(ruta).name, 0, lote, ret, nod or 0, None, True)
        if sin_valor:
            avisos.append(f"{sin_valor} celda(s) de 'no distribuida' sin valor calculado (formulas sin "
                          "resultado guardado): abre el Excel, guardalo y vuelve a correr el programa")
        if refs:
            letra = lambda c: get_column_letter(c + 1) if c is not None else "-"
            return ExcelRef(str(ruta), ws.title, letra(c_lote), letra(c_nod), letra(c_ret), refs, avisos)
    if modo_lista:
        LOTE_ID = None                    # este libro no sirve: el formato sigue sin conocerse
    return None


class ExcelNoLegible(Exception):
    pass


def buscar_excel(folder):
    """El Excel llamado EXCEL_NOMBRE (en la carpeta o subcarpetas); si no esta, el primer
    .xlsx/.xlsm con una columna 'no distribuida'."""
    todos = [p for p in folder.rglob("*") if p.is_file() and not p.name.startswith("~$")
             and CARPETA_RESULTADOS not in p.parts]
    nombrados = sorted(p for p in todos if EXCEL_NOMBRE and p.stem.lower() == EXCEL_NOMBRE.lower())
    for p in nombrados:
        if p.suffix.lower() == ".xls":
            raise ExcelNoLegible(f"{p} esta en el formato antiguo .xls: abrelo en Excel y guardalo como "
                                 f"'Libro de Excel (.xlsx)'")
        if p.suffix.lower() in (".xlsx", ".xlsm"):
            x = leer_excel_referencia(p)
            if x is None:
                raise ExcelNoLegible(f"{p} no tiene una columna 'no distribuida' con numeros de lote")
            return x
    candidatos = sorted(p for p in todos if p.suffix.lower() in (".xlsx", ".xlsm")
                        and not p.name.lower().startswith("totales_lotes"))
    for p in candidatos:
        try:
            x = leer_excel_referencia(p)
        except Exception:
            continue
        if x:
            return x
    return None


# ---------------------------- conciliacion ---------------------------------
@dataclass
class Resultado:
    ret: dict
    nod: dict
    refs: dict
    nod_fuente: dict
    problemas: list
    total: int


FUERA_EXCEL = "el lote no esta en el Excel de referencia"
ESPERADOS = {"ret": TOTAL_RETIRADO_ESPERADO, "nod": TOTAL_NO_DISTRIBUIDA_ESPERADO, "total": TOTAL_ESPERADO}


def comparar(valor, esperado):
    if esperado is None:
        return ""
    d = valor - esperado
    return "COINCIDE" if d == 0 else f"DIFERENCIA {d:+,} ({'faltan' if d < 0 else 'sobran'} {fmt(abs(d))} piezas)"
_EXCEL = None                               # ExcelRef de la corrida actual


def conciliar(lecturas):
    excel = _EXCEL
    ret, nod_doc = defaultdict(int), defaultdict(int)
    for L in lecturas.values():
        for e in L.entradas:
            if excel and SOLO_LOTES_DEL_EXCEL and e.tipo == "ret" and e.lote not in excel.refs \
                    and not e.excluido:
                e.excluido = FUERA_EXCEL
            if not e.excluido:
                (ret if e.tipo == "ret" else nod_doc)[e.lote] += e.piezas
    refs = {}
    for L in lecturas.values():
        for r in L.referencias:
            if r.lote not in refs or (r.cuadra and not refs[r.lote].cuadra):
                refs[r.lote] = r
    if excel:                               # el Excel manda sobre los reportes en PDF
        refs.update(excel.refs)
        nod = {l: r.nod for l, r in excel.refs.items() if r.nod}
        nod_fuente = {l: "Excel" for l in nod}
    else:
        nod, nod_fuente = dict(nod_doc), {l: "documentos" for l in nod_doc}
        for lote, r in refs.items():
            if r.nod is not None and lote not in nod_doc:
                nod[lote], nod_fuente[lote] = r.nod, "reporte"
    for lote, r in refs.items():
        if USAR_REPORTE_PARA_RETIRADO and r.ret is not None and lote not in ret:
            ret[lote] = r.ret
    problemas = []
    for lote in sorted(set(ret) | set(nod) | set(refs), key=orden_lote):
        rv, nv, r = ret.get(lote, 0), nod.get(lote, 0), refs.get(lote)
        if rv > MAX_RETIRADO_POR_LOTE:
            problemas.append(("tope", "ret", lote, rv, MAX_RETIRADO_POR_LOTE))
        if nv > MAX_NO_DISTRIBUIDA_POR_LOTE:
            problemas.append(("tope", "nod", lote, nv, MAX_NO_DISTRIBUIDA_POR_LOTE))
        if r and r.ret is not None and not r.acepta_ret(rv):
            problemas.append(("reporte", "ret", lote, rv, r.ret))
        if not excel and r and r.nod is not None and lote in nod_doc and nod_doc[lote] != r.nod:
            problemas.append(("reporte", "nod", lote, nod_doc[lote], r.nod))
    total = sum(ret.values()) + sum(nod.values())
    return Resultado(dict(ret), nod, refs, nod_fuente, problemas, total)


def lote_cuadra(lecturas, tipo, lote, objetivo, tope, ref=None):
    s = sum(e.piezas for L in lecturas.values() for e in L.entradas
            if not e.excluido and e.tipo == tipo and e.lote == lote)
    if ref is not None and tipo == "ret":
        return ref.acepta_ret(s)
    return s == objetivo if objetivo is not None else s <= tope


COPIA_RE = re.compile(r"copia|copy|duplicad|\(\d+\)|_\d+$|-\s*\d+$", re.I)


def ordenar_copias(entradas):
    """El original primero: los nombres con 'copia', '(1)', '_2'... van despues."""
    return sorted(entradas, key=lambda e: (bool(COPIA_RE.search(pathlib.Path(e.archivo).stem)),
                                           len(e.archivo), e.archivo))


def lote_parecido(a, b):
    """Mismo largo y un digito distinto, o dos digitos vecinos cambiados de lugar."""
    if len(a) != len(b) or a == b:
        return False
    dif = [i for i in range(len(a)) if a[i] != b[i]]
    return len(dif) == 1 or (len(dif) == 2 and dif[1] == dif[0] + 1
                             and a[dif[0]] == b[dif[1]] and a[dif[1]] == b[dif[0]])


def cambiar_lectura(lecturas, alternativas, f, alt):
    alternativas[f] = [x for x in alternativas[f] if x is not alt] + [lecturas[f]]
    lecturas[f] = alt


def valor_global():
    """(funcion, esperado) para comparar el resultado global: el total retirado si se
    conoce (es lo que sale de los PDFs); si no, el total para destruccion."""
    if ESPERADOS.get("ret") is not None:
        return (lambda r: sum(r.ret.values())), ESPERADOS["ret"]
    if ESPERADOS.get("total") is not None:
        return (lambda r: r.total), ESPERADOS["total"]
    return None


def intentar_arreglos(lecturas, alternativas, res, total_esperado, malos, correcciones):
    """Busca una explicacion exacta y unica para cada problema: otra lectura del mismo
    PDF, o un renglon que es copia exacta de otro PDF. Solo entonces se corrige."""
    for clase, tipo, lote, valor, ref in res.problemas:
        objetivo = ref if clase == "reporte" else None
        tope = ref if clase == "tope" else None
        refobj = res.refs.get(lote) if clase == "reporte" else None
        cuadra = lambda lect: lote_cuadra(lect, tipo, lote, objetivo, tope, refobj)
        if cuadra(lecturas):
            continue
        archivos = sorted({e.archivo for L in lecturas.values() for e in L.entradas
                           if not e.excluido and e.tipo == tipo and e.lote == lote})
        # a) otra lectura de un PDF
        arreglos = []
        for f in archivos:
            for alt in alternativas.get(f, []):
                prueba = dict(lecturas)
                prueba[f] = alt
                if cuadra(prueba):
                    arreglos.append((f, alt))
        if len({f for f, _ in arreglos}) == 1:
            f, alt = max(arreglos, key=lambda x: calidad(x[1]))
            cambiar_lectura(lecturas, alternativas, f, alt)
            correcciones.append(f"lote {lote}: se usa la lectura {alt.pasada} de {f} "
                                f"(con ella el lote cuadra)")
            continue
        # b) renglon copiado exactamente de otro PDF
        copias = []
        por_clave = defaultdict(list)
        for L in lecturas.values():
            for e in L.entradas:
                if not e.excluido and e.tipo == tipo and e.lote == lote:
                    por_clave[e.clave()].append(e)
        for clave, es in por_clave.items():
            if len({e.archivo for e in es}) < 2:
                continue
            for e in ordenar_copias(es)[1:]:
                e.excluido = "prueba"
                if cuadra(lecturas):
                    copias.append(e)
                e.excluido = ""
        if len({c.archivo for c in copias}) == 1:
            c = copias[0]
            orig = ordenar_copias(por_clave[c.clave()])[0]
            c.excluido = f"copia de {orig.archivo} (mismo lote y cantidad)"
            malos.setdefault(c.archivo, []).append(
                f"lote {lote}: {fmt(c.piezas)} piezas repetidas de {orig.archivo}; sin esta copia el lote "
                + (f"cuadra con {refobj.archivo}" if objetivo is not None else "queda dentro del tope"))
            continue
        if tipo != "ret" or objetivo is None:
            continue
        actual = sum(e.piezas for L in lecturas.values() for e in L.entradas
                     if not e.excluido and e.tipo == tipo and e.lote == lote)
        # d0) PDF con dos paginas iguales que no coinciden: contar la otra copia del lote
        arreglos = []
        for L in lecturas.values():
            otras = [e for e in L.entradas if e.lote == lote and e.tipo == tipo
                     and e.excluido.startswith("la pagina") and "repite los mismos lotes" in e.excluido]
            if not otras:
                continue
            contadas = [e for e in L.entradas if e.lote == lote and e.tipo == tipo and not e.excluido]
            nuevo = actual - sum(e.piezas for e in contadas) + sum(e.piezas for e in otras)
            if nuevo != actual and refobj.acepta_ret(nuevo):
                arreglos.append((L, contadas, otras))
        if len(arreglos) == 1:
            L, contadas, otras = arreglos[0]
            pag = otras[0].pagina
            for e in contadas:
                e.excluido = f"se uso la copia de la pagina {pag}, que cuadra con {refobj.archivo}"
            for e in otras:
                e.excluido = ""
            correcciones.append(f"lote {lote}: {L.archivo} tiene dos copias que no coinciden; se usa la de la "
                                f"pagina {pag} (con ella el lote cuadra exactamente)")
            continue
        # e) faltan piezas: un renglon con el numero de lote mal leido (un digito distinto
        #    o dos digitos cambiados de lugar) que no esta en el Excel y completa el lote exacto
        if actual < objetivo:
            cand = [e for L in lecturas.values() for e in L.entradas
                    if e.excluido == FUERA_EXCEL and e.tipo == tipo and lote_parecido(e.lote, lote)
                    and refobj.acepta_ret(actual + e.piezas)]
            if len(cand) == 1:
                e = cand[0]
                e.notas.append(f"CORREGIDO: el lote se leyo {e.lote}; es {lote} (un digito distinto) y con "
                               f"{fmt(e.piezas)} piezas el lote cuadra exactamente con {refobj.archivo}")
                correcciones.append(f"lote {lote}: {e.archivo} p{e.pagina} tenia el lote leido como {e.lote}")
                e.lote, e.excluido = lote, ""
                continue
        # f) sobran piezas: una raya o palomita a mano leida como digito (1, 7 o 4)
        if actual > objetivo:
            cand = [(e, v) for L in lecturas.values() for e in L.entradas
                    if not e.excluido and e.tipo == tipo and e.lote == lote and not e.verificado
                    for v in sin_marca(e.piezas) if refobj.acepta_ret(actual - e.piezas + v)]
            if len(cand) == 1:
                e, v = cand[0]
                e.notas.append(f"CORREGIDO: se leyo {fmt(e.piezas)}; una marca a mano junto al numero se leyo "
                               f"como digito; con {fmt(v)} el lote cuadra exactamente con {refobj.archivo}")
                correcciones.append(f"lote {lote}: {e.archivo} p{e.pagina} {fmt(e.piezas)} -> {fmt(v)}")
                e.piezas = v

    medir = valor_global()
    if medir is None:
        return
    total_esperado = medir[1]
    res2 = conciliar(lecturas)
    dif = medir[0](res2) - total_esperado
    if dif == 0:
        return
    # c) total global: otra lectura de un solo PDF que haga coincidir el total exacto
    arreglos = []
    for f, alts in alternativas.items():
        for alt in alts:
            prueba = dict(lecturas)
            prueba[f] = alt
            if medir[0](conciliar(prueba)) == total_esperado:
                arreglos.append((f, alt))
    if len({f for f, _ in arreglos}) == 1:
        f, alt = max(arreglos, key=lambda x: calidad(x[1]))
        cambiar_lectura(lecturas, alternativas, f, alt)
        correcciones.append(f"se usa la lectura {alt.pasada} de {f}: con ella el total coincide "
                            f"exactamente con el esperado ({fmt(total_esperado)})")
        return
    # d) total global: una copia exacta entre PDFs que explique toda la diferencia
    por_clave = defaultdict(list)
    for L in lecturas.values():
        for e in L.entradas:
            if not e.excluido:
                por_clave[e.clave()].append(e)
    candidatos = [ordenar_copias(es)[1:] for es in por_clave.values() if len({e.archivo for e in es}) > 1]
    candidatos = [e for grupo in candidatos for e in grupo if e.piezas == dif and e.lote not in res2.refs]
    if len(candidatos) == 1:
        c = candidatos[0]
        orig = ordenar_copias(por_clave[c.clave()])[0]
        c.excluido = f"copia de {orig.archivo} (mismo lote y cantidad)"
        malos.setdefault(c.archivo, []).append(
            f"lote {c.lote}: {fmt(c.piezas)} piezas repetidas de {orig.archivo}; sin esta copia el total "
            f"coincide exactamente con el esperado ({fmt(total_esperado)})")


# ------------------------- lectura en paralelo y cache -----------------------
def huella_pdf(ruta):
    """Identifica el contenido del PDF (no su nombre): renombrar un PDF no invalida la
    cache, y cambiar su contenido si."""
    h = hashlib.sha1()
    with open(ruta, "rb") as f:
        for bloque in iter(lambda: f.read(1 << 20), b""):
            h.update(bloque)
    h.update(f"|{CACHE_VERSION}|{OCR_DPI}|{OCR_DPI_REINTENTO}|{LADO_MAX_PX}|{LADO_MAX_PX_AMPLIADA}".encode())
    return h.hexdigest()[:24]


def memoria_libre_gb():
    try:
        if os.name == "nt":
            import ctypes

            class Estado(ctypes.Structure):
                _fields_ = [("dwLength", ctypes.c_ulong), ("dwMemoryLoad", ctypes.c_ulong)] + [
                    (n, ctypes.c_ulonglong) for n in ("total", "libre", "tpag", "lpag", "tvir", "lvir", "ext")]
            e = Estado()
            e.dwLength = ctypes.sizeof(Estado)
            return e.libre / 2 ** 30 if ctypes.windll.kernel32.GlobalMemoryStatusEx(ctypes.byref(e)) else None
        with open("/proc/meminfo") as f:
            for linea in f:
                if linea.startswith("MemAvailable:"):
                    return int(linea.split()[1]) / 2 ** 20
    except Exception:
        pass
    try:
        return os.sysconf("SC_AVPHYS_PAGES") * os.sysconf("SC_PAGE_SIZE") / 2 ** 30
    except Exception:
        return None


def procesos_automaticos():
    """Un proceso por nucleo (maximo 8), sin pasar de unos 1.2 GB de memoria libre por
    proceso (cada uno carga su propio OCR)."""
    n = min(os.cpu_count() or 1, 8)
    mem = memoria_libre_gb()
    if mem is not None:
        n = min(n, max(1, int(mem / 1.2)))
    return max(1, n)


_doc_abierto = None                     # (ruta, documento pdfplumber) de este proceso


def _iniciar_proceso(hilos):
    global _hilos_ocr
    _hilos_ocr = hilos
    os.environ["OMP_NUM_THREADS"] = str(hilos)
    try:
        import cv2
        cv2.setNumThreads(hilos)
    except Exception:
        pass


def _cerrar_doc():
    global _doc_abierto
    if _doc_abierto is not None:
        try:
            _doc_abierto[1].close()
        except Exception:
            pass
        _doc_abierto = None


def _leer_pagina(ruta, pnum, pasada):
    """Trabajo de un proceso: palabras de una pagina. Devuelve (tokens, fuente, segundos)."""
    global _doc_abierto
    t0 = time.perf_counter()
    if _doc_abierto is None or _doc_abierto[0] != str(ruta):
        _cerrar_doc()
        _doc_abierto = (str(ruta), pdfplumber.open(ruta))
    page = _doc_abierto[1].pages[pnum - 1]
    try:
        toks, fuente = page_tokens(page, pasada)
    finally:
        if hasattr(page, "close"):
            page.close()                # libera la memoria de la pagina
    return [tuple(t) for t in toks], fuente, time.perf_counter() - t0


class LectorPaginas:
    """Lee paginas en varios procesos a la vez y guarda lo leido en una cache en disco.

    trabajos: [(nombre, ruta, huella, pagina, lectura)]. Devuelve
    {(nombre, pagina, lectura): (tokens, fuente) o la excepcion de esa pagina}."""

    def __init__(self, carpeta_cache, procesos):
        self.cache = carpeta_cache
        if self.cache is not None:
            self.cache.mkdir(parents=True, exist_ok=True)
        self.procesos = procesos or procesos_automaticos()
        self.hilos = max(1, (os.cpu_count() or 1) // self.procesos)
        self.pool = None
        self.npags = {}
        self.de_cache = 0
        self.leidas = 0
        self.segundos = 0.0

    # ---- cache ----
    def _archivo(self, huella, pnum, pasada):
        return self.cache / f"{huella}_p{pnum:04d}_l{pasada}.json"

    def _de_cache(self, huella, pnum, pasada):
        if self.cache is None:
            return None
        for k in (0, pasada):           # 0: pagina con texto (igual en todas las lecturas)
            f = self._archivo(huella, pnum, k)
            try:
                d = json.loads(f.read_text(encoding="utf-8"))
                return [Tok(*t) for t in d["toks"]], d["fuente"]
            except FileNotFoundError:
                continue
            except Exception:           # archivo de cache danado: se vuelve a leer
                continue
        return None

    def _guardar(self, huella, pnum, pasada, toks, fuente):
        if self.cache is None:
            return
        f = self._archivo(huella, pnum, 0 if fuente == "texto" else pasada)
        tmp = f.with_suffix(".tmp")
        try:
            tmp.write_text(json.dumps({"fuente": fuente, "toks": [list(t) for t in toks]},
                                      ensure_ascii=False), encoding="utf-8")
            os.replace(tmp, f)
        except OSError:
            pass

    def limpiar(self, huellas_vigentes):
        """Borra de la cache los PDFs que ya no estan o que cambiaron."""
        if self.cache is None:
            return
        for f in self.cache.glob("*"):
            if f.name.split("_")[0] not in huellas_vigentes:
                try:
                    f.unlink()
                except OSError:
                    pass

    # ---- lectura ----
    def _pool(self):
        if self.pool is None:
            import multiprocessing
            from concurrent.futures import ProcessPoolExecutor
            # "spawn" en todos los sistemas (es lo unico que hay en Windows): un proceso nuevo
            # no hereda los hilos del OCR del proceso principal, que con "fork" se bloquean.
            self.pool = ProcessPoolExecutor(max_workers=self.procesos, initializer=_iniciar_proceso,
                                            initargs=(self.hilos,),
                                            mp_context=multiprocessing.get_context("spawn"))
        return self.pool

    def cerrar(self, rapido=False):
        if self.pool is not None:
            if rapido:
                self.pool.shutdown(wait=False, cancel_futures=True)
            else:
                self.pool.shutdown(wait=True)
            self.pool = None
        _cerrar_doc()

    def leer(self, trabajos, al_terminar=None, res=None):
        res = {} if res is None else res
        pendientes = defaultdict(int)       # nombre -> paginas que faltan
        segundos = defaultdict(float)
        leidas = defaultdict(int)           # nombre -> paginas leidas (no de la cache)
        for nombre, *_ in trabajos:
            pendientes[nombre] += 1
        faltan = []

        def listo(t, valor, seg, de_cache=False):
            nombre, _, huella, pnum, pasada = t
            res[(nombre, pnum, pasada)] = valor
            segundos[nombre] += seg
            leidas[nombre] += not de_cache
            pendientes[nombre] -= 1
            if pendientes[nombre] == 0 and al_terminar:
                al_terminar(nombre, segundos[nombre] if leidas[nombre] else None)

        for t in trabajos:
            hit = self._de_cache(t[2], t[3], t[4])
            if hit is not None:
                self.de_cache += 1
                listo(t, hit, 0.0, True)
            else:
                faltan.append(t)
        if not faltan:
            return res
        if self.procesos > 1 and len(faltan) > 1:
            faltan = self._en_paralelo(faltan, listo)
            if faltan:
                print(f"   (los procesos paralelos fallaron; se leen {len(faltan)} pagina(s) una por una)")
                self.procesos = 1           # el resto de la corrida, en un solo proceso
        for t in faltan:                    # un solo proceso
            nombre, ruta, huella, pnum, pasada = t
            try:
                toks, fuente, seg = _leer_pagina(ruta, pnum, pasada)
            except Exception as e:
                listo(t, e, 0.0)
                continue
            toks = [Tok(*x) for x in toks]
            self._guardar(huella, pnum, pasada, toks, fuente)
            self.leidas += 1
            self.segundos += seg
            listo(t, (toks, fuente), seg)
        _cerrar_doc()
        return res

    def _matar_pool(self):
        """Termina los procesos (p. ej. uno atascado) sin esperar a que acaben."""
        if self.pool is None:
            return
        for p in list(getattr(self.pool, "_processes", {}).values()):
            try:
                p.terminate()
            except Exception:
                pass
        try:
            self.pool.shutdown(wait=False, cancel_futures=True)
        except Exception:
            pass
        self.pool = None

    def _en_paralelo(self, faltan, listo):
        """Devuelve los trabajos que quedan sin hacer si los procesos fallan."""
        from concurrent.futures import FIRST_COMPLETED, wait
        from concurrent.futures.process import BrokenProcessPool
        cola = list(faltan)
        while cola:
            try:
                pool = self._pool()
                futuros = {pool.submit(_leer_pagina, t[1], t[3], t[4]): t for t in cola}
            except Exception:
                self._matar_pool()
                return cola
            cola = []
            hechos, inicio, avisado = set(), {}, {}
            pendientes = set(futuros)
            roto = atasco = False
            while pendientes:
                terminados, pendientes = wait(pendientes, timeout=5, return_when=FIRST_COMPLETED)
                ahora = time.monotonic()
                for fu in terminados:
                    t = futuros[fu]
                    try:
                        toks, fuente, seg = fu.result()
                    except BrokenProcessPool:
                        roto = True
                        continue
                    except Exception as e:
                        hechos.add(fu)
                        listo(t, e, 0.0)
                        continue
                    toks = [Tok(*x) for x in toks]
                    self._guardar(t[2], t[3], t[4], toks, fuente)
                    self.leidas += 1
                    self.segundos += seg
                    hechos.add(fu)
                    listo(t, (toks, fuente), seg)
                if roto:
                    break
                for fu in pendientes:           # paginas lentas: aviso y, si se atascan, se dejan
                    if not fu.running():
                        continue
                    inicio.setdefault(fu, ahora)
                    dur = ahora - inicio[fu]
                    nombre, _, _, pnum, pasada = futuros[fu]
                    donde = (f"{nombre}, pagina {pnum} de {self.npags.get(nombre, '?')}"
                             f"{'' if pasada == 1 else f' (lectura {pasada})'}")
                    if dur >= LIMITE_PAGINA_S:
                        print(f"   *** {donde}: mas de {LIMITE_PAGINA_S} s; se deja sin leer y sigue con lo demas")
                        hechos.add(fu)
                        listo(futuros[fu], TimeoutError(f"la pagina {pnum} tardo mas de {LIMITE_PAGINA_S} s "
                                                        f"y no se leyo"), dur)
                        atasco = True
                    elif dur >= AVISO_LENTO_S * (avisado.get(fu, 0) + 1):
                        avisado[fu] = avisado.get(fu, 0) + 1
                        print(f"   ... sigue leyendo {donde}: {int(dur)} s")
                if atasco:
                    break
            if roto:
                self._matar_pool()
                return [t for fu, t in futuros.items() if fu not in hechos]
            if atasco:                          # procesos nuevos para el resto
                self._matar_pool()
                cola = [t for fu, t in futuros.items() if fu not in hechos]
        return []


# ------------------------------- salidas -----------------------------------
ROJO = PatternFill("solid", start_color="FFC7CE", end_color="FFC7CE")


def escribir_excel(ruta, res, total_esperado):
    wb = Workbook()
    ws = wb.active
    tx = TEXTOS_EXCEL
    ws.title = tx["hoja"]
    negrita = Font(bold=True)
    lotes = sorted(set(res.ret) | set(res.nod) | set(res.refs), key=orden_lote)
    malos = {(t, l) for _, t, l, _, _ in res.problemas}
    n = len(lotes)
    fin = 3 + n                        # fila del TOTAL
    for t, titulo in enumerate(tx["titulos"]):
        c0 = 1 + 3 * t
        ws.cell(1, c0, titulo).font = negrita
        ws.cell(2, c0, tx["lote"]).font = negrita
        ws.cell(2, c0 + 1, tx["cantidad"]).font = negrita
        for i, lote in enumerate(lotes):
            fila = 3 + i
            ws.cell(fila, c0, int(lote) if lote.isdigit() else lote)
            if t == 0:
                celda = ws.cell(fila, c0 + 1, res.ret.get(lote, 0))
                if ("ret", lote) in malos:
                    celda.fill = ROJO
            elif t == 1:
                celda = ws.cell(fila, c0 + 1, res.nod.get(lote, 0))
                if ("nod", lote) in malos:
                    celda.fill = ROJO
            else:
                celda = ws.cell(fila, c0 + 1, f"=B{fila}+E{fila}")
            celda.number_format = "#,##0"
        ws.cell(fin, c0, tx["total"]).font = negrita
        col = "BEH"[t]
        celda = ws.cell(fin, c0 + 1, f"=SUM({col}3:{col}{fin - 1})" if n else 0)
        celda.number_format = "#,##0"
        celda.font = negrita
        ws.column_dimensions["ADG"[t]].width = 14
        ws.column_dimensions["BEH"[t]].width = 20
    ws.column_dimensions["C"].width = 3
    ws.column_dimensions["F"].width = 3
    valores = {"ret": sum(res.ret.values()), "nod": sum(res.nod.values()), "total": res.total}
    for t, clave in enumerate(("ret", "nod", "total")):
        e = ESPERADOS[clave] if clave != "total" else total_esperado
        if e is None:
            continue
        c0, col = 1 + 3 * t, "BEH"[t]
        ws.cell(fin + 2, c0, tx["esperado"])
        c = ws.cell(fin + 2, c0 + 1, e)
        c.number_format = "#,##0"
        ws.cell(fin + 3, c0, tx["diferencia"])
        c = ws.cell(fin + 3, c0 + 1, f"={col}{fin}-{col}{fin + 2}")
        c.number_format = "#,##0"
        if valores[clave] != e:
            c.fill = ROJO
    ws.freeze_panes = "A3"
    wb.save(ruta)


def escribir_detalle(ruta, lecturas, res):
    porlote = defaultdict(list)
    for L in lecturas.values():
        for e in L.entradas:
            porlote[e.lote].append(e)
    lineas = []
    for lote in sorted(set(porlote) | set(res.refs), key=orden_lote):
        rv, nv = res.ret.get(lote, 0), res.nod.get(lote, 0)
        lineas.append(f"{lote}   retirado {fmt(rv)}   no distribuida {fmt(nv)}   total {fmt(rv + nv)}")
        r0 = res.refs.get(lote)
        if r0 and r0.pagina == 0:
            lineas.append(f"      Excel ({r0.archivo}): no distribuida {fmt(r0.nod or 0)}, retirado segun clientes "
                          + (fmt(r0.ret) if r0.ret is not None else "-"))
        for e in sorted(porlote[lote], key=lambda e: (e.tipo, e.archivo, e.pagina)):
            estado = f"NO SUMADO: {e.excluido}" if e.excluido else (
                "verificado" if e.verificado else f"sin verificar (confianza {e.conf:.2f})")
            lineas.append(f"      {TIPOS[e.tipo]:<15} {fmt(e.piezas):>9}   {e.archivo} p{e.pagina}  [{e.forma}; {estado}]")
            for x in e.verificado + e.fallas + e.notas:
                lineas.append(f"            - {x}")
        r = res.refs.get(lote)
        if r and r.pagina:
            lineas.append(f"      reporte ({r.archivo}): retirado {fmt(r.ret) if r.ret is not None else '-'}, "
                          f"no distribuida {fmt(r.nod) if r.nod is not None else '-'}, "
                          f"total {fmt(r.total) if r.total is not None else '-'}"
                          + ("" if r.cuadra else "   (sus columnas no cuadran entre si)"))
            if res.nod_fuente.get(lote) in ("reporte", "Excel"):
                lineas.append(f"      la cantidad no distribuida se toma de {r.archivo}")
        lineas.append("")
    pathlib.Path(ruta).write_text("\n".join(lineas), encoding="utf-8")


def escribir_revision(ruta, folder, lecturas, res, total_esperado, omitidos, sin_resultado, malos,
                      correcciones, excel, info_omitidos, imagenes=None, pedidos=None):
    """Archivo para la persona que revisa: todo lo que no cuadra o no se pudo leer, con el
    motivo, lo que si se vio en cada PDF y que hacer."""
    import datetime
    t = []
    raya = "=" * 78
    activas = lambda L: [e for e in L.entradas if not e.excluido]
    imagenes = imagenes or {}
    img = lambda f, p: f"   imagen: {imagenes[(f, p)]}" if (f, p) in imagenes else ""
    ret_t, nod_t = sum(res.ret.values()), sum(res.nod.values())
    sin_res = dict(sin_resultado)
    t += [raya, f"REVISION A MANO - lotes_auto.py {VERSION} - {datetime.datetime.now():%d/%m/%Y %H:%M}",
          f"Carpeta: {folder}", raya, ""]

    # 1. Resumen
    n_pdfs = len(lecturas) + len(omitidos) + sum(1 for f in sin_res if f not in lecturas)
    con_cant = sum(1 for L in lecturas.values() if activas(L))
    reportes = sum(1 for L in lecturas.values() if L.referencias and not activas(L))
    t.append("1. RESUMEN")
    t.append(f"   PDFs encontrados: {n_pdfs}   (con cantidades: {con_cant}; reportes: {reportes}; "
             f"sin cantidades legibles: {len(sin_res)}; omitidos por FALTA REVISION: {len(omitidos)})")
    if excel:
        sk = sum(r.ret or 0 for r in excel.refs.values())
        sj = sum(r.nod or 0 for r in excel.refs.values())
        t.append(f"   Excel de referencia: {excel.ruta}")
        t.append(f"      hoja '{excel.hoja}'; columnas: lote {excel.col_lote}, no distribuida {excel.col_nod}, "
                 f"retirado segun clientes {excel.col_ret}")
        t.append(f"      {len(excel.refs)} lotes; no distribuida {fmt(sj)}; retirado segun clientes {fmt(sk)}; "
                 f"suma de ambos {fmt(sj + sk)}")
        t += [f"      AVISO: {a}" for a in excel.avisos]
    else:
        t.append("   Excel de referencia: NO SE ENCONTRO. La cantidad no distribuida sale de los PDFs y no se")
        t.append("      puede comparar lote por lote. Pon el Excel en la carpeta o usa --excel RUTA.")
    filas_res = [("Cantidad de producto retirado (PDFs)", ret_t, "ret"),
                 (f"Cantidad no distribuida ({'Excel' if excel else 'PDFs'})", nod_t, "nod"),
                 ("Cantidad total para destruccion", res.total, "total")]
    for nombre, v, k in filas_res:
        e = ESPERADOS[k]
        t.append(f"   {nombre + ':':<42}{fmt(v):>10}" + (f"   esperado {fmt(e):>10}   {comparar(v, e)}"
                                                         if e is not None else ""))
    if excel:
        t.append(f"   (Excel: retirado segun clientes {fmt(sk)}, diferencia con los PDFs {ret_t - sk:+,})")
        if ESPERADOS["nod"] is not None and sj != ESPERADOS["nod"]:
            t.append(f"   AVISO: la columna 'no distribuida' del Excel suma {fmt(sj)}, no {fmt(ESPERADOS['nod'])}: "
                     "revisa que sea el Excel correcto y que tenga todos los lotes")
    t.append("")

    # Datos por lote
    entradas_lote = defaultdict(list)
    for L in lecturas.values():
        for e in L.entradas:
            entradas_lote[(e.tipo, e.lote)].append(e)
    vistos_sin_leer = defaultdict(list)         # lote -> PDFs sin cantidades que lo mencionan
    for f, L in lecturas.items():
        if not activas(L) and not L.referencias:
            for l in sorted(set(x for linea in L.lineas for x in lotes_en(linea))):
                vistos_sin_leer[l].append(f)
    vistos_omitidos = defaultdict(list)
    for f, L in info_omitidos.items():
        for l, v in sorted({(e.lote, e.piezas) for e in L.entradas}):
            vistos_omitidos[l].append(f"{f} ({fmt(v)})")
        for l in sorted(set(x for linea in L.lineas for x in lotes_en(linea))):
            if not any(s.startswith(f) for s in vistos_omitidos[l]):
                vistos_omitidos[l].append(f)

    # 2. Que hacer
    acciones = []
    if not excel:
        acciones.append("Poner el Excel de referencia en la carpeta (o usar --excel RUTA) y volver a correr.")
    if sin_res:
        acciones.append(f"Capturar a mano los {len(sin_res)} PDF(s) de la seccion 6: abrir cada uno y anotar "
                        "lote y piezas.")
    if res.problemas:
        acciones.append(f"Revisar los {len(res.problemas)} lote(s) de la seccion 3: la seccion dice que PDFs "
                        "suman cada lote y que PDFs sin leer u omitidos lo mencionan.")
    if malos:
        acciones.append(f"Revisar los {len(malos)} PDF(s) marcados como malos (seccion 7).")
    fuera = [e for L in lecturas.values() for e in L.entradas if e.excluido == FUERA_EXCEL]
    if fuera:
        acciones.append(f"Revisar {len(fuera)} cantidad(es) con un lote que no esta en el Excel (seccion 5): "
                        "casi siempre es un numero de lote mal leido.")
    n_sosp = 0
    for L in lecturas.values():
        if L.archivo in malos:
            continue
        if L.avisos or any(e.fallas or (not e.verificado and e.conf < CONFIANZA_MINIMA)
                           or any("dudosa" in x for x in e.notas) for e in activas(L)):
            n_sosp += 1
    if n_sosp:
        acciones.append(f"{'Opcional: ' if not acciones else ''}comparar con el PDF las cantidades de los "
                        f"{n_sosp} PDF(s) sospechosos (seccion 9); se sumaron, pero no estan verificadas.")
    t.append("2. QUE HACER (en este orden)")
    t += [f"   {i}. {a}" for i, a in enumerate(acciones, 1)] or ["   Nada: todo cuadra."]
    t.append("")

    # 3. Lotes que no cuadran
    t.append(f"3. LOTES QUE NO CUADRAN: {len(res.problemas)}")
    for clase, tipo, lote, valor, ref in res.problemas:
        if clase == "tope":
            t.append(f"   lote {lote}: {TIPOS[tipo]} {fmt(valor)} supera el tope por lote de {fmt(ref)} "
                     f"(sobran al menos {fmt(valor - ref)})")
        else:
            r = res.refs.get(lote)
            d = valor - ref
            t.append(f"   lote {lote}: {TIPOS[tipo]} segun los PDFs {fmt(valor)}; segun {r.archivo if r else 'la referencia'} "
                     f"{fmt(ref)}; diferencia {d:+,} ({'faltan' if d < 0 else 'sobran'} {fmt(abs(d))} piezas)")
        es = entradas_lote.get((tipo, lote), [])
        if es:
            t.append("      lo que se encontro en los PDFs:")
            for e in sorted(es, key=lambda e: (e.archivo, e.pagina)):
                estado = (f"NO SUMADO ({e.excluido})" if e.excluido else
                          "verificado" if e.verificado else f"sin verificar, confianza OCR {e.conf:.2f}")
                t.append(f"         {fmt(e.piezas):>9}  {e.archivo}  pagina {e.pagina}  [{estado}]"
                         + img(e.archivo, e.pagina))
        else:
            t.append("      ningun PDF leido tiene este lote")
        if vistos_sin_leer.get(lote):
            t.append(f"      lo mencionan PDFs sin cantidades legibles: {', '.join(vistos_sin_leer[lote])}")
        if vistos_omitidos.get(lote):
            t.append(f"      lo mencionan PDFs omitidos por FALTA REVISION: {', '.join(vistos_omitidos[lote])}")
        if clase == "reporte":
            exactos = [f"{f} ({fmt(e.piezas)})" for f, L in info_omitidos.items() for e in L.entradas
                       if e.lote == lote and not e.excluido and e.piezas == abs(valor - ref)]
            if exactos and valor < ref:
                t.append(f"      >> la diferencia es EXACTAMENTE lo que trae: {', '.join(exactos)}")
        if clase == "reporte" and valor < ref:
            t.append("      causa probable: falta un PDF, un PDF no se pudo leer o esta marcado FALTA REVISION")
        elif clase == "reporte" or clase == "tope":
            t.append("      causa probable: un documento duplicado con otro nombre o un numero mal leido")
    t.append("")

    # 4. Lotes del Excel sin PDFs
    if excel:
        sin_pdf = [l for l, r in excel.refs.items() if (r.ret or 0) > 0 and not res.ret.get(l)]
        t.append(f"4. LOTES DEL EXCEL CON RETIRADO PERO SIN NINGUN PDF: {len(sin_pdf)}")
        for l in sorted(sin_pdf, key=orden_lote):
            extra = []
            if vistos_sin_leer.get(l):
                extra.append(f"aparece en PDFs sin leer: {', '.join(vistos_sin_leer[l])}")
            if vistos_omitidos.get(l):
                extra.append(f"aparece en omitidos: {', '.join(vistos_omitidos[l])}")
            t.append(f"   lote {l}: Excel {fmt(excel.refs[l].ret)}" + (f"   ({'; '.join(extra)})" if extra else ""))
        t.append("")

    # 5. Lotes que no estan en el Excel
    t.append(f"5. CANTIDADES CON UN LOTE QUE NO ESTA EN EL EXCEL (no sumadas): {len(fuera)}")
    for e in sorted(fuera, key=lambda e: (e.archivo, e.pagina)):
        parecidos = [l for l in (excel.refs if excel else {}) if lote_parecido(e.lote, l)]
        t.append(f"   {e.archivo}  pagina {e.pagina}: lote {e.lote}  {fmt(e.piezas)} piezas"
                 + (f"   -> se parece a {', '.join(parecidos)} del Excel" if parecidos else "")
                 + img(e.archivo, e.pagina))
    t.append("")

    # 6. PDFs sin cantidades legibles
    t.append(f"6. PDFs SIN CANTIDADES LEGIBLES: {len(sin_res)}   (texto completo en sin_resultado.log)")
    for i, (f, motivo) in enumerate(sorted(sin_res.items()), 1):
        L = lecturas.get(f)
        t.append(f"   [{i}] {f}")
        ims = [imagenes[(f, p)] for p in range(1, 4) if (f, p) in imagenes]
        if ims:
            t.append(f"       imagenes: {', '.join(ims)}")
        if L is None:
            t.append(f"       motivo: {motivo}")
            t.append("       que hacer: " + ("instalar el OCR y volver a correr, o capturar a mano"
                                             if "OCR" in motivo else "abrir el PDF; si abre, capturar a mano"))
            continue
        t.append(f"       paginas: {L.paginas}; lectura: {L.fuente}"
                 + (f" (confianza media del OCR {L.conf_media:.2f})" if L.fuente != "texto" else ""))
        lotes = sorted(set(x for linea in L.lineas for x in lotes_en(linea)))
        casi = sorted(set(m for linea in L.lineas for m in (CASI_LOTE_RE.findall(linea) if CASI_LOTE_RE else [])
                          if not lotes_en(m)))
        if not L.lineas:
            why = "no se pudo leer ningun texto (pagina en blanco o imagen ilegible)"
        elif L.fuente != "texto" and L.conf_media < 0.80:
            why = ("el texto es muy dificil de leer: casi seguro esta escrito a mano o el escaneo es de "
                   "baja calidad")
        elif not lotes:
            why = "no hay ningun numero de lote valido (8 digitos que empiecen con 25 o 26)"
        else:
            why = "se encontraron lotes pero ninguna columna de cantidades confiable"
        t.append(f"       motivo: {why}")
        t.append(f"       lotes vistos: {', '.join(lotes) if lotes else 'ninguno'}")
        if casi:
            t.append(f"       numeros que podrian ser lotes mal leidos: {', '.join(casi[:10])}")
        decis = [x for x in L.traza if "tabla" in x or "etiquetas" in x or "lado" in x]
        t += [f"       analisis: {x}" for x in decis[:4]]
        t.append("       primeras lineas leidas:")
        t += [f"          | {x}" for x in L.lineas[:6]] or ["          | (nada)"]
        t.append("       que hacer: abrir el PDF y capturar a mano lote y piezas; si es un escaneo malo, "
                 "volver a escanearlo")
    t.append("")

    # 7-8. Malos y correcciones
    t.append(f"7. PDFs MARCADOS COMO MALOS (prueba concluyente): {len(malos)}")
    for f, motivos in sorted(malos.items()):
        t.append(f"   {f}")
        t += [f"      - {m}" for m in motivos]
        ims = sorted(v for (a, p), v in imagenes.items() if a == f)
        if ims:
            t.append(f"      imagenes: {', '.join(ims)}")
    t.append("")
    t.append("8. CORRECCIONES APLICADAS (cada una tiene una prueba exacta)")
    corr = list(correcciones)
    for L in lecturas.values():
        for e in L.entradas:
            corr += [f"{e.archivo} pagina {e.pagina}: {n}" + img(e.archivo, e.pagina)
                     for n in e.notas if n.startswith("CORREGIDO")]
    t += [f"   {c}" for c in dict.fromkeys(corr)] or ["   ninguna"]
    t.append("")

    # 9. Sospechosos
    sospechosos = []
    for L in lecturas.values():
        motivos = []
        for e in activas(L):
            motivos += [f"lote {e.lote}: {x}" for x in e.fallas]
            if not e.verificado and e.conf < CONFIANZA_MINIMA:
                motivos.append(f"lote {e.lote}: {fmt(e.piezas)} leido con confianza baja ({e.conf:.2f}); "
                               "comparar con el PDF")
            motivos += [f"lote {e.lote}: {x}" for x in e.notas if "dudosa" in x or "no reconocido" in x]
        motivos += L.avisos
        if motivos and L.archivo not in malos:
            sospechosos.append((L.archivo, motivos))
    t.append(f"9. PDFs SOSPECHOSOS (se sumaron; comparar con el PDF): {len(sospechosos)}")
    for f, motivos in sospechosos:
        t.append(f"   {f}")
        t += [f"      - {m}" for m in dict.fromkeys(motivos)]
        ims = sorted(v for (a, p), v in imagenes.items() if a == f)
        if ims:
            t.append(f"      imagenes: {', '.join(ims)}")
    t.append("")

    # 10-11. Repeticiones y duplicados
    rep = [f"   {e.archivo} pagina {e.pagina}: lote {e.lote} {fmt(e.piezas)} -> {e.excluido}"
           for L in lecturas.values() for e in L.entradas
           if e.excluido and not e.excluido.startswith("copia de") and e.excluido != FUERA_EXCEL]
    t.append(f"10. REPETICIONES DENTRO DEL MISMO ARCHIVO (contadas una sola vez): {len(rep)}")
    t += rep
    t.append("")
    por_clave = defaultdict(set)
    for L in lecturas.values():
        for e in activas(L):
            por_clave[e.clave()].add(e.archivo)
    dup = [(k, fs) for k, fs in por_clave.items() if len(fs) > 1]
    t.append(f"11. MISMO LOTE Y CANTIDAD EN VARIOS ARCHIVOS (se sumaron; puede ser una copia): {len(dup)}")
    for (tipo, lote, v), fs in dup:
        t.append(f"   lote {lote} {fmt(v)}: {', '.join(sorted(fs))}")
    t.append("")

    # 12. Omitidos
    t.append(f"12. PDFs OMITIDOS POR 'FALTA REVISION' (no cuentan en ningun total): {len(omitidos)}")
    for f in omitidos:
        L = info_omitidos.get(f)
        if L is None:
            t.append(f"   {f}   (no se pudo leer)")
            continue
        vals = ", ".join(f"{e.lote}: {fmt(e.piezas)}" for e in L.entradas if not e.excluido)
        t.append(f"   {f}   (solo informativo: {vals or 'sin cantidades legibles'})")
    t.append("")

    # 13. Todos los PDFs
    t.append("13. TODOS LOS PDFs")
    filas = []
    for f in sorted(set(lecturas) | set(sin_res) | set(omitidos)):
        L = lecturas.get(f)
        if f in omitidos:
            estado = "omitido (FALTA REVISION)"
        elif L is None:
            estado = f"no se pudo leer: {sin_res.get(f, '')}"
        else:
            a = activas(L)
            r = sum(e.piezas for e in a if e.tipo == "ret")
            nd = sum(e.piezas for e in a if e.tipo == "nod")
            partes = []
            if r:
                partes.append(f"{sum(1 for e in a if e.tipo == 'ret')} cantidad(es), {fmt(r)} piezas retiradas")
            if nd:
                partes.append(f"{fmt(nd)} piezas no distribuidas"
                              + (" (solo informativo: se usa el Excel)" if excel else ""))
            if L.referencias:
                partes.append(f"reporte con {len(L.referencias)} lotes")
            if not partes:
                partes.append("SIN CANTIDADES")
            if f in malos:
                partes.append("MALO")
            estado = "; ".join(partes) + f" [{L.fuente}, {L.paginas} pag.]"
        filas.append(f"   {f}: {estado}")
    t += filas
    t.append("")
    t.append(f"14. IMAGENES DE LAS PAGINAS A REVISAR (carpeta paginas): {len(imagenes)}")
    for (f, p), nombre in sorted(imagenes.items(), key=lambda x: x[1]):
        t.append(f"   {nombre}   <- {f} pagina {p}: {', '.join((pedidos or {}).get((f, p), []))}")
    pathlib.Path(ruta).write_text("\n".join(t) + "\n", encoding="utf-8")


def paginas_a_revisar(lecturas, res, sin_resultado, malos):
    """(archivo, pagina) -> motivos, para guardar una imagen de cada pagina a revisar."""
    pedidos = defaultdict(list)
    sin_res = dict(sin_resultado)
    for f in sin_res:
        L = lecturas.get(f)
        for pag in range(1, min(L.paginas if L else 3, 3) + 1):
            pedidos[(f, pag)].append("sin cantidades legibles")
    problemas = {(t, l) for _, t, l, _, _ in res.problemas}
    for L in lecturas.values():
        for e in L.entradas:
            k = (e.archivo, e.pagina)
            if not e.excluido and (e.tipo, e.lote) in problemas:
                pedidos[k].append(f"lote {e.lote} no cuadra")
            if e.excluido == FUERA_EXCEL:
                pedidos[k].append(f"lote {e.lote} no esta en el Excel")
            if any(n.startswith("CORREGIDO") for n in e.notas):
                pedidos[k].append(f"lote {e.lote} corregido")
            if not e.excluido and (e.fallas or any("dudosa" in n for n in e.notas)
                                   or (not e.verificado and e.conf < CONFIANZA_MINIMA)):
                pedidos[k].append(f"lote {e.lote} sospechoso")
            if e.archivo in malos and (e.fallas or e.excluido.startswith("copia de")):
                pedidos[k].append("PDF marcado como malo")
    return {k: list(dict.fromkeys(v)) for k, v in pedidos.items()}


def exportar_paginas(carpeta, rutas, pedidos, dpi=110):
    """Guarda un PNG de cada pagina pedida. Devuelve (archivo, pagina) -> ruta relativa."""
    carpeta.mkdir(exist_ok=True)
    for viejo in carpeta.glob("*.png"):
        viejo.unlink()
    hechas = {}
    abiertos = {}
    try:
        for n, (f, pag) in enumerate(sorted(pedidos), 1):
            ruta = rutas.get(f)
            if ruta is None:
                continue
            try:
                if f not in abiertos:
                    abiertos[f] = pdfplumber.open(ruta)
                doc = abiertos[f]
                if pag > len(doc.pages):
                    continue
                base = re.sub(r"[^\w\- ]", "_", pathlib.Path(f).stem)
                nombre = f"{n:03d}_{base}_p{pag}.png"
                renderizar(doc.pages[pag - 1], dpi, 2200).save(carpeta / nombre)
                hechas[(f, pag)] = f"{carpeta.name}/{nombre}"
            except ImportError:
                return hechas               # sin pypdfium2 no se pueden crear imagenes
            except Exception:
                continue
    finally:
        for d in abiertos.values():
            d.close()
    return hechas


def ocr_install_hint():
    pkgs = "rapidocr onnxruntime" if sys.version_info >= (3, 13) else "rapidocr_onnxruntime"
    if os.name == "nt":
        pkgs += " msvc-runtime"
    return f'"{sys.executable}" -m pip install {pkgs}'


def resumen_archivo(L):
    ret = sum(1 for e in L.entradas if not e.excluido and e.tipo == "ret")
    nod = sum(1 for e in L.entradas if not e.excluido and e.tipo == "nod")
    partes = []
    if ret:
        partes.append(f"{ret} retirado")
    if nod:
        partes.append(f"{nod} no distribuida")
    if L.referencias:
        partes.append(f"reporte con {len(L.referencias)} lotes")
    return (", ".join(partes) or "0 cantidades") + f" ({L.fuente})"


# ---------------------------- programa principal ---------------------------
def main(argv=None):
    ap = argparse.ArgumentParser(
        prog="lotes_auto.py",
        description="Calcula por lote la cantidad retirada, la no distribuida y el total para "
                    "destruccion a partir de los PDFs de una carpeta (incluye subcarpetas).")
    ap.add_argument("carpeta", help="carpeta con los PDFs")
    ap.add_argument("--total-esperado", type=int, default=None, metavar="N",
                    help=f"total esperado para destruccion (por omision: {TOTAL_ESPERADO}; 0 = no comprobar)")
    ap.add_argument("--retirado-esperado", type=int, default=None, metavar="N",
                    help=f"total esperado de producto retirado (por omision: {TOTAL_RETIRADO_ESPERADO}; 0 = no comprobar)")
    ap.add_argument("--no-distribuida-esperada", type=int, default=None, metavar="N",
                    help=f"total esperado no distribuido (por omision: {TOTAL_NO_DISTRIBUIDA_ESPERADO}; 0 = no comprobar)")
    ap.add_argument("--excel", metavar="RUTA", default=EXCEL_REFERENCIA,
                    help="Excel de referencia (por omision se busca en la carpeta y subcarpetas)")
    ap.add_argument("--procesos", type=int, default=None, metavar="N",
                    help="procesos que leen paginas a la vez (por omision: automatico; 1 = uno solo)")
    ap.add_argument("--sin-cache", action="store_true",
                    help="volver a leer todas las paginas aunque esten en resultados_lotes/cache")
    ap.add_argument("--perfil", choices=list(PERFILES), default=None,
                    help=f"pais o idioma de los documentos (por omision: {PERFIL_PREDETERMINADO}, este retiro)")
    ap.add_argument("--sector", choices=list(SECTORES), default=None,
                    help="sector de los productos (unidades y etiquetas de lote propias del sector)")
    ap.add_argument("--config", metavar="RUTA", default=None,
                    help=f"archivo de configuracion (por omision: {CONFIG_NOMBRE} en la carpeta, si existe)")
    ap.add_argument("--version", action="version", version=f"%(prog)s {VERSION}")
    args = ap.parse_args(argv)
    inicio = time.monotonic()

    folder = pathlib.Path(args.carpeta).expanduser()
    if not folder.is_dir():
        print(f"*** No existe la carpeta: {folder}")
        return 2
    ruta_cfg = pathlib.Path(args.config).expanduser() if args.config else folder / CONFIG_NOMBRE
    try:
        config = leer_config(ruta_cfg) if (args.config or ruta_cfg.is_file()) else None
        desc = aplicar_perfil(args.perfil, args.sector, config)
    except ConfigNoValida as e:
        print(f"*** Configuracion no valida: {e}")
        return 2
    if config is not None:
        print(f"Configuracion: {ruta_cfg}")
    for d in desc:
        print(f"Perfil: {d}")
    elegir = lambda arg, valor: (arg if arg is not None else valor) or None
    esperado = elegir(args.total_esperado, TOTAL_ESPERADO)
    global ESPERADOS
    ESPERADOS = {"ret": elegir(args.retirado_esperado, TOTAL_RETIRADO_ESPERADO),
                 "nod": elegir(args.no_distribuida_esperada, TOTAL_NO_DISTRIBUIDA_ESPERADO),
                 "total": esperado}
    global _EXCEL
    if args.excel:
        ruta_x = pathlib.Path(args.excel).expanduser()
        try:
            _EXCEL = leer_excel_referencia(ruta_x) if ruta_x.is_file() else None
        except Exception as e:
            print(f"*** No se pudo leer el Excel {ruta_x}: {e}")
            return 2
        if _EXCEL is None:
            print(f"*** {ruta_x} no existe o no tiene una columna 'no distribuida' con lotes")
            return 2
    else:
        try:
            _EXCEL = buscar_excel(folder)
        except ExcelNoLegible as e:
            print(f"*** {e}")
            return 2
    if _EXCEL:
        sk = sum(r.ret or 0 for r in _EXCEL.refs.values())
        sj = sum(r.nod or 0 for r in _EXCEL.refs.values())
        print(f"Excel de referencia: {_EXCEL.ruta} (hoja '{_EXCEL.hoja}', {len(_EXCEL.refs)} lotes; "
              f"no distribuida {fmt(sj)}, retirado segun clientes {fmt(sk)})")
        for a in _EXCEL.avisos:
            print(f"   AVISO: {a}")
    else:
        print("*** No se encontro el Excel de referencia: la cantidad no distribuida saldra de los PDFs.")
    if LOTE_ID is None:
        print("*** Falta el formato del numero de lote. Indicalo en " + CONFIG_NOMBRE + ' con "batch_pattern" '
              '(por ejemplo "batch_pattern": "L\\\\d{6}"), elige un --sector que lo defina, o pon en la carpeta '
              "un Excel de referencia con una columna de lotes (Lote, Lot, Batch, Charge).")
        return 2
    salida = folder / CARPETA_RESULTADOS
    pdfs = sorted(p for p in folder.rglob("*") if p.suffix.lower() == ".pdf"
                  and salida not in p.parents)
    print(f"{len(pdfs)} PDFs encontrados en {folder} (incluye subcarpetas)")
    info_omitidos = {}

    lecturas, rutas, omitidos, sin_resultado, log = {}, {}, [], [], []
    procesos = PROCESOS if args.procesos is None else max(1, args.procesos)
    usar_cache = USAR_CACHE and not args.sin_cache
    lector = LectorPaginas(salida / "cache" if usar_cache else None, procesos)
    nombres = {str(p.relative_to(folder)): p for p in pdfs}
    huellas, npags, resultado = {}, {}, {}
    trabajos = []
    for name, pdf in nombres.items():
        try:
            huellas[name] = huella_pdf(pdf)
            with pdfplumber.open(pdf) as doc:
                npags[name] = len(doc.pages)
        except Exception as e:
            resultado[name] = ("error", e)
            print(f"{name}  -> " + ("omitido (FALTA REVISION)" if OMITIR_NOMBRE.search(pdf.name)
                                    else f"ERROR: {e}"))
            continue
        trabajos += [(name, pdf, huellas[name], p, 1) for p in range(1, npags[name] + 1)]
    lector.npags = npags
    print(f"{len(trabajos)} paginas; {lector.procesos} proceso(s) a la vez"
          + ("" if usar_cache else "; sin cache"))
    hechos = [0]
    aviso_ocr = [False]
    paginas_1 = {}

    def terminar(name, seg, pasada=1):
        hechos[0] += 1
        pdf = nombres[name]
        pre = f"[{hechos[0]}/{len(pdfs)}] {name}"
        paginas = [paginas_1[(name, p, 1)] for p in range(1, npags[name] + 1)]
        err = next((x for x in paginas if isinstance(x, Exception)), None)
        dur = f"  [{npags[name]} pag, " + ("cache]" if seg is None else f"{seg:.1f} s]")
        if OMITIR_NOMBRE.search(pdf.name):
            print(pre + "  -> omitido (FALTA REVISION; se lee solo para informar)" + dur)
            L = None
            if err is None:
                try:
                    L = analizar_documento(name, paginas, 1)
                except Exception:
                    pass
            resultado[name] = ("omitido", L)
            return
        if isinstance(err, OCRNoDisponible):
            print(pre + "  -> escaneado, OCR no disponible")
            if not aviso_ocr[0]:
                print("     Para leer PDFs escaneados instala el OCR:  " + ocr_install_hint())
                aviso_ocr[0] = True
            resultado[name] = ("sin_ocr", err)
            return
        try:
            if err is not None:
                raise err
            L = analizar_documento(name, paginas, 1)
        except Exception as e:
            print(f"{pre}  -> ERROR: {e}")
            resultado[name] = ("error", e)
            return
        resultado[name] = ("ok", L)
        print(f"{pre}  -> {resumen_archivo(L)}{dur}")

    for name in [n for n in nombres if npags.get(n) == 0]:
        terminar(name, 0.0)
    try:
        lector.leer(trabajos, terminar, paginas_1)
    except BaseException:
        lector.cerrar(rapido=True)
        raise
    # mismo orden que antes (orden de los nombres), sin importar que PDF termino primero
    for name in nombres:
        tipo, x = resultado[name]
        if tipo == "omitido" or (tipo == "error" and OMITIR_NOMBRE.search(nombres[name].name)):
            omitidos.append(name)
            if tipo == "omitido" and x is not None:
                info_omitidos[name] = x
        elif tipo == "sin_ocr":
            sin_resultado.append((name, "escaneado; OCR no disponible en este sistema"))
            log += ["=" * 70, name, "Motivo: escaneado; OCR no disponible", f"Detalle: {x}", ""]
        elif tipo == "error":
            sin_resultado.append((name, f"error al leer: {x}"))
            log += ["=" * 70, name, f"Motivo: error al leer: {x}", ""]
        else:
            lecturas[name], rutas[name] = x, nombres[name]

    # ----- comprobaciones y nuevas lecturas de las paginas dudosas -----
    res = conciliar(lecturas)
    alternativas = {}
    problema_lotes = {l for _, _, l, _, _ in res.problemas}
    medir = valor_global()
    hay_que_reintentar = bool(res.problemas) or (medir is not None and medir[0](res) != medir[1])
    dudosas = [f for f, L in lecturas.items() if L.fuente in ("OCR", "mixto") and (
        dudosa(L) or any(e.lote in problema_lotes for e in L.entradas)
        or (hay_que_reintentar and any(not e.verificado for e in L.entradas)))]
    if dudosas and (hay_que_reintentar or any(dudosa(lecturas[f]) for f in dudosas)):
        objetivo = {f: paginas_dudosas(lecturas[f], problema_lotes) for f in dudosas}
        trabajos = [(f, rutas[f], huellas[f], p, k) for f in dudosas for k in (2, 3) for p in objetivo[f]]
        print(f"\nNuevas lecturas de {len(trabajos) // 2} pagina(s) dudosa(s) de {len(dudosas)} PDF(s): "
              f"{OCR_DPI_REINTENTO} ppp con mas contraste, y ampliada al doble y enfocada...")
        paginas_r, elegidas = {}, {}
        hechos[0] = 0

        def terminar_r(f, seg):
            hechos[0] += 1
            pags = objetivo[f]
            lecturas_f = [lecturas[f]]
            errores = []
            for pasada in (2, 3) if pags else ():
                mezcla = [paginas_r[(f, p, pasada)] if p in pags else paginas_1[(f, p, 1)]
                          for p in range(1, npags[f] + 1)]
                err = next((x for x in mezcla if isinstance(x, Exception)), None)
                try:
                    if err is not None:
                        raise err
                    lecturas_f.append(analizar_documento(f, mezcla, pasada, releidas=pags))
                except Exception as e:
                    errores.append(f"  (lectura {pasada}: {e})")
            mejor = max(lecturas_f, key=calidad)
            elegidas[f] = (mejor, [x for x in lecturas_f if x is not mejor])
            if not pags:
                cuales = "ninguna pagina escaneada que releer"
            elif len(pags) < npags[f]:
                cuales = f"pagina{'s' if len(pags) > 1 else ''} {', '.join(map(str, pags))} de {npags[f]}"
            else:
                cuales = f"{npags[f]} pag"
            cuando = "" if not pags else ", cache" if seg is None else f", {seg:.1f} s"
            print(f"   [{hechos[0]}/{len(dudosas)}] {f} ({cuales}{cuando})"
                  f"{''.join(errores)}  -> se usa la lectura {mejor.pasada}: {resumen_archivo(mejor)}")

        for f in dudosas:
            if not objetivo[f]:
                terminar_r(f, 0.0)
        try:
            lector.leer(trabajos, terminar_r, paginas_r)
        except BaseException:
            lector.cerrar(rapido=True)
            raise
        for f in dudosas:                   # mismo orden que antes
            mejor, otras = elegidas[f]
            alternativas[f] = otras
            lecturas[f] = mejor
    lector.cerrar()
    if usar_cache:
        lector.limpiar(set(huellas.values()))
    malos, correcciones = {}, []
    res = conciliar(lecturas)
    if res.problemas or (medir is not None and medir[0](res) != medir[1]):
        intentar_arreglos(lecturas, alternativas, res, esperado, malos, correcciones)
        res = conciliar(lecturas)
    for f, alts in alternativas.items():    # las lecturas no coinciden
        for alt in alts:
            otra = defaultdict(int)
            for e in alt.entradas:
                if not e.excluido:
                    otra[(e.tipo, e.lote)] += e.piezas
            for e in lecturas[f].entradas:
                if not e.excluido and (e.tipo, e.lote) in otra and otra[(e.tipo, e.lote)] != e.piezas:
                    e.notas.append(f"dudosa: la lectura {alt.pasada} dio {fmt(otra[(e.tipo, e.lote)])}")
    for L in lecturas.values():             # lo que coincide con el reporte queda verificado
        for e in L.entradas:
            r = res.refs.get(e.lote)
            ref = None if r is None else (r.ret if e.tipo == "ret" else r.nod)
            total = res.ret.get(e.lote) if e.tipo == "ret" else res.nod.get(e.lote)
            if not e.excluido and ref is not None and (ref == total or (e.tipo == "ret" and r.acepta_ret(total))):
                e.verificado.append(f"el total del lote coincide con {r.archivo}")
    for f, L in lecturas.items():           # cifras que contradicen el propio documento
        fallas = sorted({x for e in L.entradas if not e.excluido for x in e.fallas})
        if fallas:
            alts = alternativas.get(f, [])
            if all(any(e.fallas for e in alt.entradas if not e.excluido) for alt in alts):
                malos.setdefault(f, []).extend(f"sus propias cifras no cuadran: {x}" for x in fallas)

    for f, L in lecturas.items():
        if not L.entradas and not L.referencias:
            sin_resultado.append((f, "sin cantidades legibles"))
            log += ["=" * 70, f, f"Motivo: sin cantidades legibles (fuente: {L.fuente})", "Texto leido:"]
            log += ["  | " + x for x in L.lineas] or ["  | (nada: no se pudo leer texto)"]
            log += ["Analisis:"] + ["  - " + x for x in L.traza] + [""]

    # ----- salidas -----
    salida.mkdir(exist_ok=True)
    out = salida / "totales_lotes.xlsx"
    try:
        escribir_excel(out, res, esperado)
    except PermissionError:
        print(f"\n*** No se pudo guardar {out.name}: esta abierto en Excel. Cierralo y vuelve a correr el script.")
        return 1
    escribir_detalle(salida / "detalle_lotes.txt", lecturas, res)
    todas_rutas = {str(p.relative_to(folder)): p for p in pdfs}
    pedidos = paginas_a_revisar(lecturas, res, sin_resultado, malos)
    imagenes = exportar_paginas(salida / "paginas", todas_rutas, pedidos) if pedidos else {}
    if pedidos and not imagenes:
        print("(no se pudieron crear las imagenes de las paginas a revisar en este sistema)")
    rev = salida / "revisar_a_mano.txt"
    escribir_revision(rev, folder, lecturas, res, esperado, omitidos, sin_resultado, malos, correcciones,
                      _EXCEL, info_omitidos, imagenes, pedidos)
    logf = salida / "sin_resultado.log"
    if log:
        logf.write_text("\n".join(log) + "\n", encoding="utf-8")
    elif logf.exists():
        logf.unlink()

    ret_t, nod_t = sum(res.ret.values()), sum(res.nod.values())
    print()
    for nombre, v, k in (("Cantidad de producto retirado", ret_t, "ret"),
                         ("Cantidad no distribuida", nod_t, "nod"),
                         ("Cantidad total para destruccion", res.total, "total")):
        e = ESPERADOS[k]
        print(f"{nombre + ':':<33}{fmt(v):>12}" + (f"   esperado {fmt(e):>10}   "
              + ("COINCIDE" if v == e else "*** " + comparar(v, e)) if e is not None else ""))
    if res.problemas:
        print(f"*** {len(res.problemas)} lote(s) no cuadran (tope o reporte); ver revisar_a_mano.txt")
    if malos:
        print(f"*** {len(malos)} PDF(s) marcados como MALOS:")
        for f, motivos in sorted(malos.items()):
            print(f"      {f}: {motivos[0]}")
    if correcciones:
        print(f"{len(correcciones)} correccion(es) aplicadas; ver revisar_a_mano.txt")
    print(f"{len(omitidos)} omitidos por FALTA REVISION | {len(sin_resultado)} sin cantidades legibles")
    print(f"Revision: {rev}")
    print(f"Detalle por lote: {salida / 'detalle_lotes.txt'}")
    if imagenes:
        print(f"Imagenes de {len(imagenes)} pagina(s) a revisar: {salida / 'paginas'}")
    print(f"Todos los resultados estan en: {salida}")
    print(f"Excel: {out}")
    print(f"Tiempo: {time.monotonic() - inicio:.0f} s ({lector.leidas} pagina(s) leidas, "
          f"{lector.de_cache} de la cache, {lector.procesos} proceso(s))")
    return 0


if __name__ == "__main__":
    sys.exit(main())
