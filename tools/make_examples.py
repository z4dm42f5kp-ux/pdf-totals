"""Generate the example folder: 120 synthetic PDF documents of a product withdrawal,
the reference Excel file and the expected results.

All documents, batch numbers, companies and quantities are invented. The layouts imitate
the kinds of documents that lotes_auto.py reads: delivery receipts with a table, return
forms with unit price and total price, documents with labelled lines, an inventory list,
a withdrawal report, merged PDFs with a repeated page, an exact copy of another PDF,
files named FALTA REVISION and scanned (image only) documents.

The generator also calculates the correct result. It writes it to expected.json, which
tools/check_examples.py compares with the result of lotes_auto.py.

Usage:
    python tools/make_examples.py [FOLDER]
    python tools/make_examples.py --check

With --check, the examples are generated into a temporary folder and compared with
examples/. The PDF files and expected.json must be identical, and the Excel file must
contain the same values. Return value 1 means that examples/ is out of date.

The default folder is examples/ in the project folder. Requires reportlab, openpyxl,
pypdfium2 and Pillow (pypdfium2 and Pillow are installed together with pdfplumber).
The output is the same on every run (fixed random seed).
"""
import datetime
import filecmp
import json
import pathlib
import tempfile
import random
import shutil
import sys

import pypdfium2 as pdfium
from openpyxl import Workbook
from PIL import Image, ImageFilter
from reportlab.lib.pagesizes import landscape, letter
from reportlab.pdfgen import canvas

ROOT = pathlib.Path(__file__).resolve().parent.parent
SEED = 20261015
# Fixed dates make the files identical on every run.
FECHA_FIJA = datetime.datetime(2026, 10, 15, 12, 0, 0)

# 24 invented batch numbers: eight digits, starting with 25 or 26.
LOTES = [f"25{n:06d}" for n in (240101, 240117, 240123, 240139, 240142, 240158, 240166, 240171,
                               240189, 240193, 240205, 240214)] + \
        [f"26{n:06d}" for n in (240003, 240019, 240027, 240031, 240048, 240052, 240064, 240075,
                               240086, 240090, 240107, 240118)]
# Differences that lotes_auto.py is known to produce on these examples. They depend on
# the fixed seed: change them together. check_examples.py accepts exactly these.
KNOWN_DIFFERENCES = [
    {"lote": "25240123", "kind": "retirado", "difference": -954,
     "file": "098 16-JUL DEVOLUCION ESCANEADA TIJUANA.pdf",
     "reason": "the OCR does not read the pieces cell of this one-row scanned table; "
               "revisar_a_mano.txt lists the PDF as unreadable and names the batch"},
]
PRODUCTO = {"25": ("CEFTRIAXONA I.M. 1 FA 1G/3.5 ML", "COL242"),
            "26": ("TRAXICOLL IM 1 FA 1G/3.5 ML", "COL237")}
CEDIS = ["CULIACAN", "MONTERREY", "TIJUANA", "HERMOSILLO", "GUADALAJARA", "MERIDA", "PUEBLA",
         "LEON", "QUERETARO", "TOLUCA", "MEXICALI", "CHIHUAHUA"]
CLIENTES = ["DISTRIBUIDORA DEL NORTE", "FARMACIAS DEL PACIFICO", "COMERCIALIZADORA SALUD",
            "GRUPO FARMACEUTICO CENTRO", "DROGUERIA DEL BAJIO", "SUMINISTROS MEDICOS DEL SUR"]
MESES = ["JUL", "AGO", "SEP"]


def fmt(n):
    return f"{n:,}"


def grid(c, x, y, widths, heights):
    yy = y
    for h in [0] + heights:
        yy -= h
        c.line(x, yy, x + sum(widths), yy)
    xx = x
    for w in [0] + widths:
        c.line(xx, y, xx, y - sum(heights))
        xx += w


def cells(c, x, y, widths, h, values, size=7, bold=False):
    c.setFont("Helvetica-Bold" if bold else "Helvetica", size)
    xx = x
    for w, v in zip(widths, values):
        lines = v.split("\n") if v else []
        for i, line in enumerate(lines):
            yy = y - h / 2 - size / 3 + (len(lines) - 1) * size * 0.55 - i * size * 1.1
            c.drawCentredString(xx + w / 2, yy, line)
        xx += w


def acuse_page(c, filas, cedis, fecha):
    """Delivery receipt: one table with pieces, product code, name, batch, expiry, origin."""
    c.setFont("Helvetica-Bold", 20)
    c.drawString(60, 720, "Distribuidora Ejemplo")
    c.setFont("Helvetica-Bold", 13)
    c.drawString(300, 725, "ACUSE DE ENTREGA DE PRODUCTO")
    c.setFont("Helvetica-Bold", 8)
    c.drawString(60, 690, "FECHA DE ENTREGA:")
    c.drawString(60, 670, "CEDIS:")
    c.drawString(60, 650, "PROVEEDOR:")
    c.setFont("Helvetica", 8)
    c.drawString(200, 690, fecha)
    c.drawString(200, 670, cedis)
    c.drawString(160, 650, "LABORATORIO EJEMPLO S.A. DE C.V.")
    c.setFont("Helvetica-Bold", 8)
    c.drawString(390, 690, "MOTIVO DE ENTREGA:")
    c.setFont("Helvetica", 7)
    c.drawString(390, 675, "     Devolución a proveedor")
    c.drawString(390, 662, "X   Retiro de producto de mercado")
    w = [70, 65, 160, 80, 70, 70]
    grid(c, 60, 630, w, [22] + [15] * max(len(filas), 6))
    cells(c, 60, 630, w, 22, ["No. de piezas", "Clave del\nproducto", "Nombre del producto", "Lote",
                               "Fecha de\ncaducidad", "CEDIS de origen"], 7, True)
    y = 608
    for lote, piezas in filas:
        nombre, clave = PRODUCTO[lote[:2]]
        cells(c, 60, y, w, 15, [str(piezas), clave, nombre, lote, "30/MAY/2027", cedis], 7)
        y -= 15


def acuse(path, filas, cedis, fecha):
    c = canvas.Canvas(str(path), pagesize=letter, invariant=1)
    acuse_page(c, filas, cedis, fecha)
    c.save()


def devolucion(path, filas, cliente, fecha, rnd):
    """Return form: pieces, batch, unit price and total price, with printed totals."""
    c = canvas.Canvas(str(path), pagesize=landscape(letter), invariant=1)
    c.setFont("Helvetica-Bold", 12)
    c.drawString(30, 575, cliente)
    c.setFont("Helvetica-Bold", 14)
    c.drawString(300, 570, "DEVOLUCION PROVEEDORES")
    c.setFont("Helvetica", 6)
    c.drawString(70, 535, "Proveedor:")
    c.drawString(270, 535, "LABORATORIO EJEMPLO")
    c.drawString(560, 500, f"Fecha de Entrega:   {fecha}")
    w = [55, 150, 45, 55, 45, 60, 65, 40, 55, 150]
    grid(c, 18, 470, w, [24] + [16] * max(len(filas), 4))
    cells(c, 18, 470, w, 24, ["CODIGO SAP", "DESCRIPCION", "CANTIDAD EN\nPIEZAS", "LOTE", "CADUCIDAD",
                               "NO FACTURA FISCAL", "FECHA DE FACTURA", "PRECIO\nUNITARIO", "IMPORTE TOTAL",
                               "MOTIVO DE DEVOLUCION"], 5, True)
    y = 446
    total_p = total_m = 0
    for lote, piezas in filas:
        precio = rnd.choice([21, 23, 24, 26, 27, 29])
        importe = piezas * precio
        total_p += piezas
        total_m += importe
        cells(c, 18, y, w, 16, ["300483", PRODUCTO[lote[:2]][0][:30], fmt(piezas), lote, "01/07/2027",
                                 f"FE{rnd.randint(200000, 299999)}", "28/02/2026", f"${precio}.00",
                                 f"${fmt(importe)}.00", "Retiro de mercado"], 5.5)
        y -= 16
    c.setFont("Helvetica-Bold", 6)
    c.drawString(80, 352, "TOTAL DE PIEZAS")
    c.drawString(280, 352, fmt(total_p))
    c.drawString(470, 352, "MONTO TOTAL")
    c.drawString(560, 352, f"${fmt(total_m)}.00")
    c.save()


def etiquetas(path, lote, piezas, cliente, fecha):
    """Letter with labelled lines: LOTE: ... and CANTIDAD EN PIEZAS: ..."""
    c = canvas.Canvas(str(path), pagesize=letter, invariant=1)
    c.setFont("Helvetica-Bold", 14)
    c.drawString(60, 720, "CONSTANCIA DE RETIRO DE PRODUCTO")
    c.setFont("Helvetica", 10)
    lineas = [f"CLIENTE: {cliente}", f"FECHA: {fecha}", f"PRODUCTO: {PRODUCTO[lote[:2]][0]}",
              f"LOTE: {lote}", f"CANTIDAD EN PIEZAS: {fmt(piezas)}",
              "MOTIVO: Retiro de producto de mercado"]
    y = 680
    for linea in lineas:
        c.drawString(60, y, linea)
        y -= 22
    c.drawString(60, 480, "Firma de quien entrega: ______________________")
    c.save()


def inventario(path, nod):
    c = canvas.Canvas(str(path), pagesize=letter, invariant=1)
    c.setFont("Helvetica-Bold", 12)
    c.drawString(60, 730, "INVENTARIO DE PRODUCTO NO DISTRIBUIDO")
    c.setFont("Helvetica", 8)
    c.drawString(60, 712, "Almacén: CENTRAL     Fecha de corte: 30/09/2026")
    filas = [(l, q) for l, q in nod.items() if q]
    w = [70, 170, 60, 80, 60]
    grid(c, 60, 690, w, [20] + [15] * len(filas))
    cells(c, 60, 690, w, 20, ["LOTE", "DESCRIPCIÓN", "CADUCIDAD", "EXISTENCIA\n(PIEZAS)", "UBICACIÓN"],
          7, True)
    y = 670
    for i, (lote, q) in enumerate(filas):
        cells(c, 60, y, w, 15, [lote, PRODUCTO[lote[:2]][0], "jun-27", str(q), f"A-{i + 1:02d}"], 7)
        y -= 15
    c.save()


def reporte(path, filas):
    """Withdrawal report: manufactured, distributed, not distributed, withdrawn, total."""
    c = canvas.Canvas(str(path), pagesize=landscape(letter), invariant=1)
    c.setFont("Helvetica-Bold", 10)
    c.drawString(40, 570, "Relación de lotes")
    w = [25, 150, 60, 45, 55, 60, 60, 60, 60, 60]
    grid(c, 30, 540, w, [20] + [14] * len(filas))
    cells(c, 30, 540, w, 20, ["No.", "Denominación genérica", "Lote", "Cad.", "Fecha", "Dato 1", "Dato 2",
                               "Dato 3", "Dato 4", "Dato 5"], 6, True)
    y = 520
    for i, (lote, fab, dist, nod, ret) in enumerate(filas, 1):
        nombre = "TRAXICOLL (CEFTRIAXONA)" if lote.startswith("26") else "CEFTRIAXONA"
        cells(c, 30, y, w, 14, [str(i), nombre, lote, "may-27", "07-oct-25", fmt(fab), fmt(dist), fmt(nod),
                                 fmt(ret), fmt(nod + ret)], 6.5)
        y -= 14
    c.save()


def escanear(src, dst, rnd, dpi=200):
    """Turn a PDF into an image-only PDF, like a scan: slightly rotated, blurred, with noise."""
    pdf = pdfium.PdfDocument(str(src))
    paginas = []
    for page in pdf:
        img = page.render(scale=dpi / 72).to_pil().convert("L")
        img = img.rotate(rnd.uniform(-0.6, 0.6), fillcolor=255, resample=Image.BICUBIC)
        img = img.filter(ImageFilter.GaussianBlur(0.5))
        px = img.load()
        for _ in range(img.size[0] * img.size[1] // 800):
            px[rnd.randrange(img.size[0]), rnd.randrange(img.size[1])] = rnd.randint(140, 210)
        paginas.append(img)
    pdf.close()
    paginas[0].save(dst, save_all=True, append_images=paginas[1:], resolution=dpi, quality=70,
                    creationDate=FECHA_FIJA.timetuple(), modDate=FECHA_FIJA.timetuple())


def acuse_varias_paginas(path, paginas, cedis):
    """Several delivery receipts in one PDF, one page each: [(filas, fecha), ...]."""
    c = canvas.Canvas(str(path), pagesize=letter, invariant=1)
    for filas, fecha_pag in paginas:
        acuse_page(c, filas, cedis, fecha_pag)
        c.showPage()
    c.save()


def excel(path, ret, nod):
    wb = Workbook()
    ws = wb.active
    ws.title = "Retiro"
    ws["A1"] = "Relación de lotes - retiro de mercado (datos de ejemplo)"
    enc = ["No.", "Denominación", "Lote", "Caducidad", "Fecha", "Fabricado", "Liberado", "Distribuido",
           "Clientes", "Cantidad no distribuida\n(piezas)( inventarios", "Cantidad de producto\nretirado (clientes)",
           "Cantidad de producto total para destrucción"]
    for j, h in enumerate(enc, 1):
        ws.cell(3, j, h)
    for i, lote in enumerate(LOTES, 4):
        ws.cell(i, 1, i - 3)
        ws.cell(i, 2, "TRAXICOLL" if lote.startswith("26") else "CEFTRIAXONA")
        ws.cell(i, 3, int(lote))
        ws.cell(i, 4, "may-27")
        ws.cell(i, 6, 80000)
        ws.cell(i, 10, nod[lote])
        ws.cell(i, 11, ret[lote])
        ws.cell(i, 12, nod[lote] + ret[lote])
    n = 4 + len(LOTES)
    ws.cell(n, 2, "TOTAL")
    ws.cell(n, 10, sum(nod.values()))
    ws.cell(n, 11, sum(ret.values()))
    ws.cell(n, 12, sum(nod.values()) + sum(ret.values()))
    wb.properties.created = wb.properties.modified = FECHA_FIJA
    wb.save(path)


class Mundo:
    """Keeps the correct result while the documents are created."""

    def __init__(self, rnd):
        self.rnd = rnd
        self.ret = {l: 0 for l in LOTES}
        self.usados = set()
        self.docs = []

    def filas(self, n, minimo=40, maximo=2400):
        """n rows of (batch, pieces). A (batch, pieces) pair is never used twice, so that
        two different documents are never mistaken for copies of each other."""
        lotes = self.rnd.sample(LOTES, n)
        out = []
        for lote in lotes:
            while True:
                q = self.rnd.randint(minimo, maximo)
                if (lote, q) not in self.usados:
                    break
            self.usados.add((lote, q))
            out.append((lote, q))
        return out

    def cuenta(self, nombre, filas, tipo):
        for lote, q in filas:
            self.ret[lote] += q
        self.docs.append({"file": nombre, "type": tipo, "counted": True,
                          "rows": [{"lote": l, "piezas": q} for l, q in filas]})

    def no_cuenta(self, nombre, filas, tipo, motivo):
        self.docs.append({"file": nombre, "type": tipo, "counted": False, "reason": motivo,
                          "rows": [{"lote": l, "piezas": q} for l, q in filas]})


def fecha(rnd):
    return f"{rnd.randint(1, 28):02d}/{rnd.choice(MESES)}/2026"


def comprobar():
    """--check: generate into a temporary folder and compare with examples/."""
    from openpyxl import load_workbook
    actual = ROOT / "examples"
    with tempfile.TemporaryDirectory() as tmp:
        nuevo = pathlib.Path(tmp) / "examples"
        main(["make_examples.py", str(nuevo)])
        diferentes = []
        nombres = {p.name for p in nuevo.iterdir()} | {p.name for p in actual.iterdir()
                                                         if p.name not in ("README.md", "resultados_lotes")}
        for nombre in sorted(nombres):
            a, b = actual / nombre, nuevo / nombre
            if not a.exists() or not b.exists():
                diferentes.append(nombre)
            elif nombre.endswith(".xlsx"):
                va = [[c.value for c in fila] for fila in load_workbook(a).active.iter_rows()]
                vb = [[c.value for c in fila] for fila in load_workbook(b).active.iter_rows()]
                if va != vb:
                    diferentes.append(nombre)
            elif not filecmp.cmp(a, b, shallow=False):
                diferentes.append(nombre)
    if diferentes:
        print(f"examples/ is out of date ({len(diferentes)} files differ), for example:")
        for nombre in diferentes[:10]:
            print("  " + nombre)
        print("Run: python tools/make_examples.py")
        return 1
    print("examples/ is up to date.")
    return 0


def main(argv):
    if argv[1:] == ["--check"]:
        return comprobar()
    out = pathlib.Path(argv[1]) if len(argv) > 1 else ROOT / "examples"
    rnd = random.Random(SEED)
    if out.exists():
        for p in out.iterdir():
            if p.name in ("README.md",):
                continue
            shutil.rmtree(p) if p.is_dir() else p.unlink()
    out.mkdir(parents=True, exist_ok=True)
    tmp = out / ".tmp"
    tmp.mkdir()
    m = Mundo(rnd)
    n = 0

    def nombre(prefijo, cedis):
        nonlocal n
        n += 1
        return f"{n:03d} {rnd.randint(1, 28):02d}-{rnd.choice(MESES)} {prefijo} {cedis}.pdf"

    # 50 delivery receipts with text, 1 to 5 rows each.
    for _ in range(50):
        cedis = rnd.choice(CEDIS)
        f = nombre("ACUSE", cedis)
        filas = m.filas(rnd.randint(1, 5))
        acuse(out / f, filas, cedis, fecha(rnd))
        m.cuenta(f, filas, "acuse")

    # 20 return forms with unit price and total price.
    for _ in range(20):
        cliente = rnd.choice(CLIENTES)
        f = nombre("DEVOLUCION", cliente.split()[-1])
        filas = m.filas(rnd.randint(1, 4))
        devolucion(out / f, filas, cliente, fecha(rnd), rnd)
        m.cuenta(f, filas, "devolucion")

    # 12 letters with labelled lines.
    for _ in range(12):
        cliente = rnd.choice(CLIENTES)
        f = nombre("CONSTANCIA", cliente.split()[-1])
        [(lote, q)] = m.filas(1)
        etiquetas(out / f, lote, q, cliente, fecha(rnd))
        m.cuenta(f, [(lote, q)], "etiquetas")

    # 20 scanned documents (image only): 14 delivery receipts and 6 return forms.
    # Few rows and small quantities, so that the OCR has an easy job.
    for i in range(20):
        cedis = rnd.choice(CEDIS)
        filas = m.filas(rnd.randint(1, 3), 40, 1500)
        texto = tmp / f"scan{i}.pdf"
        if i < 14:
            f = nombre("ACUSE ESCANEADO", cedis)
            acuse(texto, filas, cedis, fecha(rnd))
            tipo = "acuse escaneado"
        else:
            f = nombre("DEVOLUCION ESCANEADA", cedis)
            devolucion(texto, filas, rnd.choice(CLIENTES), fecha(rnd), rnd)
            tipo = "devolucion escaneada"
        escanear(texto, out / f, rnd)
        m.cuenta(f, filas, tipo)

    # 4 merged PDFs: a delivery receipt printed twice and one other receipt. The
    # repeated page counts once.
    for i in range(4):
        cedis = rnd.choice(CEDIS)
        f = nombre("VARIOS (3 paginas)", cedis)
        a, b = m.filas(rnd.randint(2, 4)), m.filas(rnd.randint(1, 3))
        fa, fb = fecha(rnd), fecha(rnd)
        acuse_varias_paginas(out / f, [(a, fa), (b, fb), (a, fa)], cedis)
        m.cuenta(f, a + b, "unido, pagina repetida")

    # 6 files named FALTA REVISION: they are read but never counted.
    for _ in range(6):
        cedis = rnd.choice(CEDIS)
        n += 1
        f = f"{n:03d} {rnd.randint(1, 28):02d}-{rnd.choice(MESES)} ACUSE {cedis} (FALTA REVISION).pdf"
        filas = m.filas(rnd.randint(1, 3))
        acuse(out / f, filas, cedis, fecha(rnd))
        m.no_cuenta(f, filas, "acuse", "FALTA REVISION")

    # More delivery receipts with text, so that there are 120 PDFs in total
    # (three more files follow: the inventory, the copy and the report).
    while n < 117:
        cedis = rnd.choice(CEDIS)
        f = nombre("ACUSE", cedis)
        filas = m.filas(rnd.randint(1, 4))
        acuse(out / f, filas, cedis, fecha(rnd))
        m.cuenta(f, filas, "acuse")

    # The not distributed quantities, the inventory list and the withdrawal report.
    nod = {l: (rnd.choice([0, 0, rnd.randint(20, 900), rnd.randint(1000, 30000)])) for l in LOTES}
    n += 1
    f = f"{n:03d} INVENTARIO NO DISTRIBUIDO.pdf"
    inventario(out / f, nod)
    m.docs.append({"file": f, "type": "inventario", "counted": False,
                   "reason": "not distributed quantities come from the reference Excel file",
                   "rows": [{"lote": l, "piezas": q} for l, q in nod.items() if q]})

    # One exact copy of a delivery receipt under another name. It must not count twice.
    copia_de = next(d for d in m.docs if d["type"] == "acuse")
    n += 1
    f = f"{n:03d} {copia_de['file'][4:-4]} (copia).pdf"
    shutil.copyfile(out / copia_de["file"], out / f)
    m.no_cuenta(f, [(r["lote"], r["piezas"]) for r in copia_de["rows"]], "copia",
                f"exact copy of {copia_de['file']}")

    # The withdrawal report: its withdrawn column equals the correct result.
    filas_rep = []
    for lote in LOTES:
        fab = 80000 + rnd.randint(0, 9) * 1000
        filas_rep.append((lote, fab, fab - nod[lote], nod[lote], m.ret[lote]))
    n += 1
    f = f"{n:03d} RELACION DE LOTES.pdf"
    reporte(out / f, filas_rep)
    m.docs.append({"file": f, "type": "reporte", "counted": False,
                   "reason": "the withdrawal report is used to check the other documents",
                   "rows": [{"lote": r[0], "piezas": r[4]} for r in filas_rep]})

    shutil.rmtree(tmp)
    excel(out / "Retirado_retiro_producto_Traxicoll.xlsx", m.ret, nod)
    esperado = {
        "note": "Generated by tools/make_examples.py. All data is invented.",
        "retirado": m.ret,
        "no_distribuida": nod,
        "totals": {"retirado": sum(m.ret.values()), "no_distribuida": sum(nod.values()),
                   "total": sum(m.ret.values()) + sum(nod.values())},
        "known_differences": KNOWN_DIFFERENCES,
        "documents": m.docs,
    }
    (out / "expected.json").write_text(json.dumps(esperado, indent=1, ensure_ascii=False) + "\n",
                                       encoding="utf-8")
    pdfs = len(list(out.glob("*.pdf")))
    print(f"{pdfs} PDFs in {out}")
    print(f"retirado {sum(m.ret.values()):,}  no distribuida {sum(nod.values()):,}  "
          f"total {sum(m.ret.values()) + sum(nod.values()):,}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
